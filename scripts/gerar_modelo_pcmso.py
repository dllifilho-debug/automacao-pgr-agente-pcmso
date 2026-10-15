"""Modelo .docx do PCMSO (texto fixo Seconci 2026) a partir do PCMSO Mauá/Plateau D'Or — D-ARQ-96.

Entrada: o .docx convertido do RTF do acervo, que é o que o soffice produz de
`matrizes_originais/PCMSO(ATUALIZAÇÃO)MAUÁ ... 28.09.26.rtf.doc` (copiar para `.rtf` e
`soffice --headless --convert-to docx`). Saída: o modelo, com

- campos `{{NOME}}` que o app preenche, realçados em turquesa;
- correções `C-NN`, realçadas em amarelo e com comentário do Word dizendo o motivo;
- sugestões `S-NN`, só comentário (texto intacto);
- os blocos por GHE trocados pelo campo `{{BLOCOS_GHE}}`.

`--limpo` gera o modelo congelado: mesmo texto, sem realces nem comentários (só depois da
aprovação do Diovanni). A lista "o que mudou e por quê" é `docs/referencia/MODELO_PCMSO_CORRECOES.md`.

    python -m scripts.gerar_modelo_pcmso <maua.docx> [--limpo] [--saida <arquivo>]
"""

from __future__ import annotations

import argparse
import copy
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

_RAIZ = Path(__file__).resolve().parent.parent
SAIDA_REVISAO = _RAIZ / "agente_medico" / "superficie" / "modelos" / "pcmso_seconci_2026_revisao.docx"
SAIDA_LIMPO = _RAIZ / "agente_medico" / "superficie" / "modelos" / "pcmso_seconci_2026.docx"

AUTOR_COMENTARIO = "Revisão do modelo (proposta)"
_LEGIVEL = {
    "{{RAZAO_SOCIAL}}": "[nome da empresa, preenchido automaticamente]",
    "{{OBRA}}": "[nome da obra, preenchido automaticamente]",
}
_LIMITE_CITACAO = 160

CAMPOS = (
    "RAZAO_SOCIAL", "OBRA", "DATA_DOCUMENTO", "VERSAO", "ENDERECO", "COMPLEMENTO", "CNPJ",
    "CEP", "CIDADE", "BAIRRO", "UF", "CNAE", "GRAU_RISCO", "DESCRICAO_CNAE",
    "REVISAO_DATA", "REVISAO_NUMERO", "REVISAO_CONSIDERACOES",
    "CRONOGRAMA_INICIO", "CRONOGRAMA_FIM",
    *(f"CRONOGRAMA_MES_{n:02d}" for n in range(1, 14)),
    "BLOCOS_GHE",
    "MEDICO_NOME", "MEDICO_CRM", "MEDICO_NIT", "MEDICO_CPF", "MEDICO_ESPECIALIDADE", "MEDICO_RQE",
)


class ModeloInesperado(RuntimeError):
    """O documento de entrada não tem o texto que a correção espera — nada é gravado."""


@dataclass(frozen=True)
class Troca:
    """Troca de texto num parágrafo. `vezes` é o número exato de parágrafos onde `antigo` aparece."""

    antigo: str
    novo: str
    marca: str  # "C-NN", "S-NN" (só comentário) ou "CAMPO"
    motivo: str = ""
    vezes: int = 1


def campo(nome: str) -> str:
    if nome not in CAMPOS:
        raise KeyError(nome)
    return "{{" + nome + "}}"


_C01 = ("O documento trazia o nome de outra empresa (\"SPE INCORPORAÇÃO H6 LTDA\"), que sobrou de um PCMSO anterior. "
        "Agora o nome da empresa e o da obra são preenchidos automaticamente em cada PCMSO.")
_C05_SEGUINTE = ("Renumeração: havia dois itens 9.3.4; com Radiações Ionizantes passando a 9.3.5, este item e os "
                 "seguintes sobem um número.")
_C04 = ("O texto mandava consultar um anexo que não existe com esse número. O Questionário OMS de epilepsia está no "
        "Anexo 16 e o fluxograma de acidente com material biológico, no Anexo 16.4.")
_C03 = ("O texto mandava consultar os fluxos nos itens 11.2 a 11.4, mas eles estão nos anexos 16.1 a 16.3 "
        "deste documento.")
_C11 = "Redação ajustada à nota sobre PNOS do Quadro 2 do Anexo III da NR-07, conferida no texto oficial."

TROCAS: tuple[Troca, ...] = (
    # --- campos: capa, cabeçalho de página, seção 1, seção 2 ---
    Troca("MAUÁ PARTICIPAÇÕES ESTRUTURAS S/A", campo("RAZAO_SOCIAL"), "CAMPO", vezes=1),
    Troca("28/09/2026", campo("DATA_DOCUMENTO"), "CAMPO", vezes=1),
    Troca("Versão: 1", "Versão: " + campo("VERSAO"), "CAMPO"),
    Troca("MAUA PARTICIPACOES ESTRUTURADAS S/A - PLATEAU D´OR",
          campo("RAZAO_SOCIAL") + " - " + campo("OBRA"), "CAMPO"),
    Troca("Rua T 65 , 345", campo("ENDERECO"), "CAMPO"),
    Troca("02.689.014/0001-90", campo("CNPJ"), "CAMPO"),
    Troca("74823-370", campo("CEP"), "CAMPO"),
    Troca("Setor Bela Vista", campo("BAIRRO"), "CAMPO"),
    Troca("41.20-4-00", campo("CNAE"), "CAMPO"),
    Troca("Construção de Edifícios", campo("DESCRICAO_CNAE"), "CAMPO"),
    Troca("28.09.2026", campo("REVISAO_DATA"), "CAMPO"),
    Troca("Atualização do PCMSO", campo("REVISAO_CONSIDERACOES"), "CAMPO"),
    Troca("Unidade: PLATEAU D´OR", "Unidade: " + campo("OBRA"), "CAMPO"),
    # --- seção 3 ---
    Troca("Portaria SEPRT n.º 6.734, de 10/03/2020", "Portaria SEPRT n.º 6.734, de 09/03/2020", "C-14",
          "A data da portaria aparecia de duas formas: 10/03/2020 aqui e 9 de março de 2020 no parágrafo seguinte. "
          "A tabela de publicação da NR-07 registra 9 de março de 2020 (Diário Oficial de 13/03/2020). Atenção: o "
          "próprio texto oficial da NR-07 escreve \"10 de março\" em outra linha; a data a manter fica para "
          "confirmação."),
    # --- seção 6 ---
    Troca("11.2 Fluxo Geral", "16.1 Fluxo Geral", "C-03", _C03),
    Troca("11.3 Fluxo Específico", "16.2 Fluxo Específico", "C-03", _C03),
    Troca("11.4 Fluxo Específico", "16.3 Fluxo Específico", "C-03", _C03),
    Troca("PGR/PCMAT", "PGR/PCMAT", "S-03",
          "Sugestão: escrever só \"PGR\". O PCMAT foi substituído pelo PGR na NR-18 revisada e só continua valendo "
          "para obras iniciadas antes dela (item 18.17.1 da NR-18)."),
    Troca("A monitoração da exposição ocupacional a agentes químicos",
          "A monitoração da exposição ocupacional a agentes químicos", "S-04",
          "Sugestão de redação: a frase diz que a monitoração \"ficará a critério do PGR\", mas pela NR-07 quem "
          "define os exames é o médico responsável pelo PCMSO, a partir dos riscos que o PGR identifica e "
          "classifica (itens 7.5.12 e 7.5.18 da NR-07)."),
    # --- seção 8 ---
    Troca("Avaliação do estado de saúde dos empregados em atividades críticas, como definidas nesta Norma, "
          "considerando os riscos envolvidos em cada situação e a investigação de patologias que possam "
          "impedir o exercício de tais atividades com segurança.",
          "i) Avaliação do estado de saúde dos empregados em atividades críticas, como definidas nesta Norma, "
          "considerando os riscos envolvidos em cada situação e a investigação de patologias que possam "
          "impedir o exercício de tais atividades com segurança.", "C-13",
          "As responsabilidades do médico responsável iam de a) a h), seguidas de um parágrafo sem letra e de um "
          "\"J)\" maiúsculo. O parágrafo sem letra passa a ser a alínea i)."),
    Troca("J) Prestar esclarecimentos", "j) Prestar esclarecimentos", "C-13",
          "Letra minúscula, na sequência das alíneas (a seguinte à nova alínea i)."),
    # --- seção 9.2 e 9.3 ---
    Troca("VI. No Exame Demissional", "V. No Exame Demissional", "C-07",
          "A numeração dos tipos de exame pulava do IV para o VI. O exame demissional é o V."),
    Troca("135 (centro e trinta e cinco)", "135 (cento e trinta e cinco)", "C-08",
          "Erro de digitação (\"centro\" em vez de \"cento\"). O erro vem do próprio texto oficial da NR-07 "
          "(item 7.5.11); aqui foi corrigido."),
    Troca("o levantamento preliminar de o PGR", "o levantamento preliminar do PGR", "C-09",
          "Erro de digitação. A NR-07 (item 7.5.12, alínea a) diz \"do PGR\"."),
    Troca("c) Os exames complementares laboratoriais previstos nesta NR devem ser executados por laboratório que "
          "atenda ao disposto na RDC/ANVISA n.º 302/2005, no que se refere aos procedimentos de coleta, "
          "acondicionamento, transporte e análise, e interpretados com base nos critérios constantes nos Anexos "
          "da Norma, NR-07, e são obrigatórios quando:",
          "c) Encaminhar o empregado à Previdência Social, quando houver afastamento do trabalho superior a 15 "
          "(quinze) dias, para avaliação de incapacidade e definição da conduta previdenciária;", "C-06",
          "Na lista do que cabe à empresa quando um exame revela doença relacionada ao trabalho, a alínea c) "
          "repetia por engano o parágrafo sobre os laboratórios. Entra o texto da NR-07 (item 7.5.19.5, alínea c): "
          "encaminhamento à Previdência Social nos afastamentos acima de 15 dias, o mesmo que o documento já traz "
          "no item 9.3.2."),
    # --- 9.3.4 a 9.3.9 ---
    Troca("9.3.4 Substâncias Química Cancerígenas", "9.3.4 Substâncias Químicas Cancerígenas", "C-15",
          "Concordância: \"Substâncias Químicas\"."),
    Troca("9.3.4- Radiações Ionizantes", "9.3.5 - Radiações Ionizantes", "C-05",
          "Havia dois itens com o número 9.3.4 (Substâncias Químicas Cancerígenas e Radiações Ionizantes). "
          "Radiações Ionizantes passa a 9.3.5, e os itens seguintes sobem um número até o 9.3.9."),
    Troca("9.3.5 - Exames para Trabalho em Altura", "9.3.6 - Exames para Trabalho em Altura", "C-05", _C05_SEGUINTE),
    Troca("trabalhador\">", "trabalhador.\"", "C-15",
          "A citação da NR-35 (item 35.4.4.1) terminava com os sinais '\">', resto de formatação. A citação foi "
          "conferida com o texto oficial da NR-35."),
    Troca("9.3.6- Exames para Trabalho em Espaço Confinado", "9.3.7 - Exames para Trabalho em Espaço Confinado",
          "C-05", _C05_SEGUINTE),
    Troca("9.3.7 Motoristas e Operadores de Máquinas Pesadas", "9.3.8 Motoristas e Operadores de Máquinas Pesadas",
          "C-05", _C05_SEGUINTE),
    Troca("9.3.8 Segurança e Saúde no Trabalho em Serviços de Saúde (NR 32)",
          "9.3.9 Segurança e Saúde no Trabalho em Serviços de Saúde (NR 32)", "C-05",
          _C05_SEGUINTE + " A seção da NR-32 foi mantida, porque o Seconci é serviço de saúde e segue a NR-32. "
          "Fica em aberto se ela deve aparecer só quando o PGR apontar risco biológico (sugestão S-01)."),
    Troca("9.3.8.1- Saúde do Trabalhador", "9.3.9.1 - Saúde do Trabalhador", "C-05", _C05_SEGUINTE),
    Troca("9.3.8.2- Cuidados", "9.3.9.2 - Cuidados", "C-05", _C05_SEGUINTE),
    Troca("9.3.8.3 -Condutas", "9.3.9.3 - Condutas", "C-05", _C05_SEGUINTE),
    Troca("9.3.8.4. Unidades", "9.3.9.4. Unidades", "C-05", _C05_SEGUINTE),
    Troca("9.3.8.5. Fluxograma", "9.3.9.5. Fluxograma", "C-05", _C05_SEGUINTE),
    Troca("(Anexo 11.1)", "(Anexo 16)", "C-04", _C04, vezes=3),
    Troca("(Anexo 11.5)", "(Anexo 16.4)", "C-04", _C04),
    Troca("independentemente do tipo do tipo de exame", "independentemente do tipo de exame", "C-15",
          "Palavras repetidas (\"do tipo do tipo\").", vezes=3),
    Troca("esclarecimento diagnostico", "esclarecimento diagnóstico", "C-15",
          "Acentuação: \"diagnóstico\".", vezes=3),
    # --- 9.4 audiometria ---
    Troca("(nível de audição em decibéis).ienos constantes das alíneas ma devem seguir o preconizado na norma "
          "técnica ISO 8253-1, e os resultados devem ser incluídos em certificado de aferição e/ou calibração que "
          "acompanhará o equipamento.",
          "(nível de audição em decibéis).", "C-10",
          "O fim deste item do Anexo II da NR-07 (audiometria em ambiente silencioso) estava corrompido: repetia, "
          "truncado, o item anterior (\"ienos constantes das alíneas ma…\"). O trecho repetido foi retirado; o "
          "restante confere com o texto oficial da NR-07."),
    # --- 9.5 RX (Anexo III, Quadro 2) ---
    Troca("**Para ser classificado como PNOS", "***Para ser classificado como PNOS", "C-11",
          "A nota sobre PNOS é marcada com três asteriscos, como no título do Quadro 2 do Anexo III da NR-07."),
    Troca("fluído pulmonar", "fluido pulmonar", "C-11", _C11),
    Troca("não se emitir radiação ionizante", "não ser emissor de radiação ionizante", "C-11", _C11),
    Troca("não causar defeitos tóxicos além da inflamação", "não causar efeitos tóxicos além de inflamação", "C-11",
          "O documento dizia \"defeitos tóxicos\", o que muda o sentido. A NR-07 (Anexo III, nota sobre PNOS) diz "
          "\"efeitos tóxicos além de inflamação\"."),
    # --- 9.9 calor ---
    Troca("Conforme o item 4.2.3 da NR 9", "Conforme o item 4.2.3 do Anexo III da NR-09", "C-12",
          "O item 4.2.3 citado é do Anexo III (Calor) da NR-09, e não do texto principal da norma."),
    Troca("item 4 do Anexo da NR-09", "item 4 do Anexo III da NR-09", "C-12",
          "O item 4 citado é do Anexo III (Calor) da NR-09."),
    # --- 10 ---
    Troca("referentes aos empegados", "referentes aos empregados", "C-15", "Erro de digitação: \"empregados\"."),
    # --- 12 ---
    Troca("referente à Lei 15337.", "referente à Lei nº 15.377/2026.", "C-02",
          "O número da lei aparecia de duas formas no documento (15.377 e 15337). O correto é Lei nº 15.377, de 2 de "
          "abril de 2026, que incluiu o art. 169-A na CLT (conferida no site do Planalto)."),
    # --- 13 cronograma ---
    Troca("JULHO DE 2026", campo("CRONOGRAMA_INICIO"), "CAMPO"),
    Troca("A JULHO DE 2027", "A " + campo("CRONOGRAMA_FIM"), "CAMPO"),
    Troca("(OBS: ADENDO NÃO PRECISA DESSE ITEM)", "(OBS: ADENDO NÃO PRECISA DESSE ITEM)", "S-06",
          "Sugestão: esta observação é uma nota interna de edição, mas sai impressa no documento. Retirar do texto; "
          "o sistema passa a gerar o adendo sem o cronograma."),
    Troca("EMPRESA: SPE INCORPORAÇÃO H6 LTDA", "EMPRESA: " + campo("RAZAO_SOCIAL"), "C-01", _C01),
    Troca("Local: SPE INCORPORAÇÃO H6 LTDA", "Local: " + campo("OBRA"), "C-01", _C01),
    # --- 15 termo ---
    Troca("Empresa SPE INCORPORAÇÃO H6 LTDA o compromisso", "Empresa " + campo("RAZAO_SOCIAL") + " o compromisso",
          "C-01", _C01),
    # --- 16 ---
    Troca("FLUXO GERAL DE SUGESTÃO DE CONDUTAS EM MEDICINA DO TRABALHO",
          "FLUXO GERAL DE SUGESTÃO DE CONDUTAS EM MEDICINA DO TRABALHO", "S-07",
          "Atenção: a figura deste anexo é a mesma do anexo 16.2 (fluxo das alterações nos exames de monitoração "
          "biológica do Anexo I da NR-07). O fluxo geral de condutas não está no documento; falta a figura correta."),
    Troca("reavaliação dos iscos ocupacionais", "reavaliação dos riscos ocupacionais", "C-15",
          "Erro de digitação: \"riscos\".", vezes=2),
    Troca("(imunoglobina)", "(imunoglobulina)", "C-15", "Termo correto: \"imunoglobulina\"."),
    Troca("Patrícia Montalvo Moraes", campo("MEDICO_NOME"), "CAMPO"),
    Troca("CRM: 14949/GO", "CRM: " + campo("MEDICO_CRM"), "CAMPO"),
    Troca("NIT(PIS/PASEP):", "NIT(PIS/PASEP): " + campo("MEDICO_NIT"), "CAMPO"),
    Troca("CPF:002.631.081-39", "CPF: " + campo("MEDICO_CPF"), "CAMPO"),
    Troca("Especialidade: Medicina do Trabalho", "Especialidade: " + campo("MEDICO_ESPECIALIDADE"), "CAMPO"),
    Troca("RQE: 9.662", "RQE: " + campo("MEDICO_RQE"), "CAMPO"),
)

# Comentários sobre títulos que ficam intactos.
SUGESTOES_EM_TITULO: tuple[Troca, ...] = (
    Troca("9.3.1 – QUADRO 1 DO ANEXO I NR-7", "", "S-02",
          "Os Quadros 1 e 2 do Anexo I da NR-07 estão colados como imagem; se a NR-07 mudar, a imagem não "
          "acompanha. Sugestão: trocar por uma referência à NR-07 vigente ou por uma tabela em texto conferida "
          "com a norma."),
    Troca("1. FUNDAMENTAÇÃO LEGAL", "", "S-09",
          "Falta o parágrafo único do art. 169-A da CLT (Lei nº 15.377/2026): a empresa deve informar aos "
          "empregados que podem deixar de comparecer ao serviço, sem prejuízo do salário, para os exames "
          "preventivos de HPV e dos cânceres de mama, de colo do útero e de próstata (inciso XII do art. 473 da "
          "CLT). Texto sugerido para o fim deste item: \"Conforme o parágrafo único do art. 169-A, a empresa "
          "informará aos empregados a possibilidade de deixar de comparecer ao serviço para a realização desses "
          "exames preventivos, sem prejuízo do salário, nos termos do inciso XII do art. 473 da CLT.\""),
    Troca("13.2 - PROGRAMAÇÃO DAS AÇÕES", "", "S-08",
          "No modelo original, o que se preenche à mão estava destacado em amarelo. Esse destaque foi retirado "
          "para não se confundir com as correções. Período, meses e nome da empresa passam a ser preenchidos "
          "automaticamente; as datas e os temas das palestras desta tabela continuam a ser preenchidos à mão."),
)

def _citacao(texto: str) -> str:
    for marcador, legivel in _LEGIVEL.items():
        texto = texto.replace(marcador, legivel)
    if len(texto) > _LIMITE_CITACAO:
        texto = texto[: _LIMITE_CITACAO].rsplit(" ", 1)[0] + " …"
    return f"“{texto}”"


def texto_comentario(troca: Troca) -> str:
    """Comentário que a médica valida sozinha: o que mudou, de quê para quê e por quê; o número do item no fim."""
    if troca.marca.startswith("S-"):
        return f"SUGESTÃO (o texto não foi alterado). {troca.motivo} (item {troca.marca} da lista de correções)"
    return (
        f"CORREÇÃO. Antes: {_citacao(troca.antigo)}. Agora: {_citacao(troca.novo)}. {troca.motivo} "
        f"(item {troca.marca} da lista de correções)"
    )


_INICIO_GHE = "Unidade: "
_INICIO_POS_GHE = "11. PROGRAMA DE VACINAÇÃO"
_MESES_CRONOGRAMA = ("JUL", "AGO", "SET", "OUT", "NOV", "DEZ", "JAN", "FEV", "MAR", "ABR", "MAI", "JUN")


def _texto(elemento: object) -> str:
    return "".join(t.text or "" for t in elemento.iter(qn("w:t")))  # type: ignore[attr-defined]


def _todos_paragrafos(doc: DocumentoDocx) -> list[Paragraph]:
    paragrafos = [Paragraph(p, doc._body) for p in doc.element.body.iter(qn("w:p"))]
    for secao in doc.sections:
        for parte in (secao.header, secao.even_page_header, secao.first_page_header):
            paragrafos.extend(Paragraph(p, parte) for p in parte._element.iter(qn("w:p")))
    return paragrafos


def _dividir(run: Run, posicao: int) -> Run:
    """Divide o run em `posicao`; devolve o run da direita (o da esquerda fica no lugar)."""
    texto = run.text
    novo = copy.deepcopy(run._r)
    run._r.addnext(novo)
    run.text = texto[:posicao]
    direita = Run(novo, run._parent)
    direita.text = texto[posicao:]
    return direita


def _isolar(paragrafo: Paragraph, inicio: int, fim: int) -> list[Run]:
    """Runs que cobrem exatamente [inicio, fim) do texto do parágrafo, dividindo os das pontas."""
    cursor = 0
    cobertos: list[Run] = []
    for run in list(paragrafo.runs):
        texto = run.text
        a, b = cursor, cursor + len(texto)
        cursor = b
        if not texto or b <= inicio or a >= fim:
            continue
        if a < inicio:
            run = _dividir(run, inicio - a)
            a = inicio
        if b > fim:
            _dividir(run, fim - a)
        cobertos.append(run)
    return cobertos


def _realcar(run: Run, marca: str) -> None:
    run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE if marca == "CAMPO" else WD_COLOR_INDEX.YELLOW


def _aplicar(doc: DocumentoDocx, paragrafos: list[Paragraph], troca: Troca, marcar: bool) -> None:
    alvos = [p for p in paragrafos if troca.antigo in "".join(r.text for r in p.runs)]
    if len(alvos) != troca.vezes:
        raise ModeloInesperado(f"{troca.marca}: {troca.antigo!r} em {len(alvos)} parágrafo(s), esperado {troca.vezes}")
    for paragrafo in alvos:
        texto = "".join(r.text for r in paragrafo.runs)
        inicio = texto.index(troca.antigo)
        runs = _isolar(paragrafo, inicio, inicio + len(troca.antigo))
        if not runs:
            raise ModeloInesperado(f"{troca.marca}: {troca.antigo!r} fora de runs de texto")
        so_comentario = troca.marca.startswith("S-")
        if not so_comentario:
            runs[0].text = troca.novo
            for sobra in runs[1:]:
                sobra._r.getparent().remove(sobra._r)
            runs = runs[:1]
        if not marcar:
            continue
        if not so_comentario:
            _realcar(runs[0], troca.marca)
        if troca.marca != "CAMPO" and _no_corpo(doc, paragrafo):
            doc.add_comment(runs, text=texto_comentario(troca), author=AUTOR_COMENTARIO, initials="RM")


def _no_corpo(doc: DocumentoDocx, paragrafo: Paragraph) -> bool:
    corpo = doc.element.body
    elemento = paragrafo._p.getparent()
    while elemento is not None:
        if elemento is corpo:
            return True
        elemento = elemento.getparent()
    return False


def _sugestao_em_titulo(doc: DocumentoDocx, paragrafos: list[Paragraph], troca: Troca) -> None:
    alvos = [
        p for p in paragrafos
        if troca.antigo in p.text and _no_corpo(doc, p) and not (p.style is not None and (p.style.name or "").upper().startswith("TOC"))
    ]
    if len(alvos) != 1:
        raise ModeloInesperado(f"{troca.marca}: {troca.antigo!r} em {len(alvos)} parágrafo(s)")
    runs = [r for r in alvos[0].runs if r.text]
    doc.add_comment(runs, text=texto_comentario(troca), author=AUTOR_COMENTARIO, initials="RM")


def _consertar_tc_literais(doc: DocumentoDocx) -> int:
    """A conversão do RTF deixa `TC "título"` como texto visível antes de um campo vazio; vira campo TC de novo."""
    consertados = 0
    for t in list(doc.element.body.iter(qn("w:t"))):
        texto = t.text or ""
        if not texto.startswith('TC "'):
            continue
        run = t.getparent()
        inicio = run.getnext()
        fld = inicio.find(qn("w:fldChar")) if inicio is not None else None
        if fld is None or fld.get(qn("w:fldCharType")) != "begin":
            raise ModeloInesperado(f"TC literal sem campo depois: {texto!r}")
        instr_run = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {texto} \\l 1 "
        instr_run.append(instr)
        inicio.addnext(instr_run)
        run.getparent().remove(run)
        consertados += 1
    return consertados


def _tirar_realce_original(doc: DocumentoDocx) -> int:
    """O modelo do Seconci marca em amarelo o que se preenche à mão; amarelo aqui passa a ser só correção C-NN."""
    removidos = 0
    for realce in list(doc.element.body.iter(qn("w:highlight"))):
        realce.getparent().remove(realce)
        removidos += 1
    return removidos


def _trocar_blocos_ghe(doc: DocumentoDocx, marcar: bool) -> int:
    filhos = list(doc.element.body.iterchildren())
    unidade = [i for i, e in enumerate(filhos) if e.tag == qn("w:tbl") and _texto(e).strip().startswith(_INICIO_GHE)]
    fim = [i for i, e in enumerate(filhos) if e.tag == qn("w:tbl") and _texto(e).strip().startswith(_INICIO_POS_GHE)]
    if len(unidade) != 1 or len(fim) != 1 or unidade[0] >= fim[0]:
        raise ModeloInesperado(f"limites dos blocos GHE: Unidade={unidade}, seção 11={fim}")
    removidos = filhos[unidade[0] + 1 : fim[0]]
    for elemento in removidos:
        doc.element.body.remove(elemento)
    p = OxmlElement("w:p")
    filhos[unidade[0]].addnext(p)
    paragrafo = Paragraph(p, doc._body)
    run = paragrafo.add_run(campo("BLOCOS_GHE"))
    if marcar:
        _realcar(run, "CAMPO")
    return sum(1 for e in removidos if e.tag == qn("w:tbl"))


def _tirar_ghe_do_indice(doc: DocumentoDocx) -> int:
    removidos = 0
    for p in list(doc.element.body.iter(qn("w:p"))):
        estilo = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if estilo is None or not (estilo.get(qn("w:val")) or "").startswith("TOC"):
            continue
        if _texto(p).startswith(("Setor: GHE", "Cargo: ")):
            p.getparent().remove(p)
            removidos += 1
    return removidos


def _campos_do_cronograma(doc: DocumentoDocx, marcar: bool) -> None:
    """Cabeçalho MÊS/ANO do 13.1: 13 colunas (mês inicial … mesmo mês do ano seguinte)."""
    for tabela in doc.element.body.iter(qn("w:tbl")):
        linhas = tabela.findall(qn("w:tr"))
        for tr in linhas:
            celulas = tr.findall(qn("w:tc"))
            textos = [_texto(c).strip() for c in celulas]
            if not textos or textos[0] != "MÊSANO":
                continue
            meses = [c for c, t in zip(celulas, textos) if t[:3] in _MESES_CRONOGRAMA]
            if len(meses) != 13:
                raise ModeloInesperado(f"cronograma com {len(meses)} meses, esperado 13")
            for numero, celula in enumerate(meses, start=1):
                paragrafos = celula.findall(qn("w:p"))
                for sobra in paragrafos[1:]:
                    celula.remove(sobra)
                primeiro = Paragraph(paragrafos[0], doc._body)
                modelo_run = next((r for r in primeiro.runs if r.text), None)
                for r in list(primeiro.runs):
                    r._r.getparent().remove(r._r)
                run = primeiro.add_run(campo(f"CRONOGRAMA_MES_{numero:02d}"))
                if modelo_run is not None and modelo_run._r.rPr is not None:
                    run._r.insert(0, copy.deepcopy(modelo_run._r.rPr))
                if marcar:
                    _realcar(run, "CAMPO")
            return
    raise ModeloInesperado("linha MÊS/ANO do cronograma não encontrada")


def _campo_em_celula_vazia(doc: DocumentoDocx, rotulo: str, nome: str, marcar: bool) -> None:
    """Célula da seção 1 com só o rótulo (Complemento): o valor entra como parágrafo novo, no estilo do rótulo."""
    for tc in doc.element.body.iter(qn("w:tc")):
        if _texto(tc).strip() != rotulo:
            continue
        rotulo_p = tc.findall(qn("w:p"))[0]
        novo = copy.deepcopy(rotulo_p)
        for filho in list(novo):
            if filho.tag != qn("w:pPr"):
                novo.remove(filho)
        rotulo_p.addnext(novo)
        run = Paragraph(novo, doc._body).add_run(campo(nome))
        if marcar:
            _realcar(run, "CAMPO")
        return
    raise ModeloInesperado(f"célula {rotulo!r} não encontrada")


def _valor_de_rotulo(doc: DocumentoDocx, rotulo: str, valor: str, nome: str, marcar: bool) -> None:
    """Valor curto e ambíguo ("3", "1", "GO") trocado só na célula logo após o rótulo."""
    for tc in doc.element.body.iter(qn("w:tc")):
        paragrafos = tc.findall(qn("w:p"))
        textos = [_texto(p).strip() for p in paragrafos]
        if rotulo in textos:
            i = textos.index(rotulo)
            seguinte = [j for j in range(i + 1, len(paragrafos)) if textos[j]]
            if seguinte and textos[seguinte[0]] == valor:
                p = Paragraph(paragrafos[seguinte[0]], doc._body)
                runs = [r for r in p.runs if r.text]
                runs[0].text = campo(nome)
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
                if marcar:
                    _realcar(runs[0], "CAMPO")
                return
    raise ModeloInesperado(f"rótulo {rotulo!r} seguido de {valor!r} não encontrado")


def _revisao_numero(doc: DocumentoDocx, marcar: bool) -> None:
    for tr in doc.element.body.iter(qn("w:tr")):
        celulas = tr.findall(qn("w:tc"))
        if len(celulas) == 3 and _texto(celulas[0]).strip() == campo("REVISAO_DATA"):
            p = Paragraph(celulas[1].findall(qn("w:p"))[0], doc._body)
            runs = [r for r in p.runs if r.text]
            if [r.text for r in runs] != ["1"]:
                raise ModeloInesperado("número da revisão diferente de '1'")
            runs[0].text = campo("REVISAO_NUMERO")
            if marcar:
                _realcar(runs[0], "CAMPO")
            return
    raise ModeloInesperado("linha do controle de revisão não encontrada")


def _pedir_atualizacao_de_campos(doc: DocumentoDocx) -> None:
    """O Word atualiza o índice ao abrir (as entradas dos GHE vêm do app)."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        elemento = OxmlElement("w:updateFields")
        elemento.set(qn("w:val"), "true")
        settings.append(elemento)


def gerar_modelo(entrada: Path, saida: Path, *, marcar: bool) -> dict[str, int]:
    doc = Document(str(entrada))
    realces = _tirar_realce_original(doc)
    tc = _consertar_tc_literais(doc)
    tabelas_ghe = _trocar_blocos_ghe(doc, marcar)
    indice = _tirar_ghe_do_indice(doc)
    paragrafos = _todos_paragrafos(doc)
    for troca in TROCAS:
        if troca.antigo in ("MAUÁ PARTICIPAÇÕES ESTRUTURAS S/A", "28/09/2026"):
            _aplicar_em_todos(doc, paragrafos, troca, marcar)
        else:
            _aplicar(doc, paragrafos, troca, marcar)
    _valor_de_rotulo(doc, "Cidade", "Goiânia", "CIDADE", marcar)
    _valor_de_rotulo(doc, "UF", "GO", "UF", marcar)
    _valor_de_rotulo(doc, "Grau de Risco", "3", "GRAU_RISCO", marcar)
    _campo_em_celula_vazia(doc, "Complemento", "COMPLEMENTO", marcar)
    _revisao_numero(doc, marcar)
    _campos_do_cronograma(doc, marcar)
    if marcar:
        for sugestao in SUGESTOES_EM_TITULO:
            _sugestao_em_titulo(doc, paragrafos, sugestao)
    _pedir_atualizacao_de_campos(doc)
    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))
    return {"realces_originais": realces, "tc_consertados": tc, "tabelas_ghe_removidas": tabelas_ghe, "entradas_indice_ghe": indice}


def _aplicar_em_todos(doc: DocumentoDocx, paragrafos: list[Paragraph], troca: Troca, marcar: bool) -> None:
    """Razão social e data: capa + cabeçalhos de página (quantos houver, pelo menos a capa)."""
    alvos = [p for p in paragrafos if troca.antigo in "".join(r.text for r in p.runs)]
    if not alvos:
        raise ModeloInesperado(f"{troca.antigo!r} não encontrado")
    for paragrafo in alvos:
        _aplicar(doc, [paragrafo], troca, marcar)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("entrada", type=Path, help=".docx convertido do RTF do PCMSO Mauá")
    parser.add_argument("--limpo", action="store_true", help="sem realces nem comentários (modelo congelado)")
    parser.add_argument("--saida", type=Path, default=None)
    args = parser.parse_args(argv)
    saida = args.saida or (SAIDA_LIMPO if args.limpo else SAIDA_REVISAO)
    try:
        contagem = gerar_modelo(args.entrada, saida, marcar=not args.limpo)
    except ModeloInesperado as erro:
        print(f"modelo inesperado: {erro}", file=sys.stderr)
        return 1
    print(f"{saida}: {contagem}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
