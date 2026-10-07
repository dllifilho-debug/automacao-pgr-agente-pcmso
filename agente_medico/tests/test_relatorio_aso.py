"""D-ARQ-91 fatia 2 — documento "riscos para o ASO" e rodapé "Uso interno —
confidencial" nos anexos das médicas. Cada teste nomeia a reversão de código que o
deixa vermelho."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from agente_medico.motor.protocolo import carregar
from agente_medico.motor.tipos import (
    Criterio,
    ExameEmitido,
    MatrizGHE,
    RelatorioASO,
    Resultado,
    SugestaoASO,
)
from agente_medico.superficie.documento_matriz import (
    CabecalhoDocumento,
    RodapeDocumento,
    montar_documento,
    renderizar_docx,
)
from agente_medico.superficie.memorial_matriz import (
    RODAPE_CONFIDENCIAL,
    montar_memorial,
    renderizar_memorial_docx,
)
from agente_medico.superficie.relatorio_aso import (
    TEXTO_INEXISTENCIA,
    montar_relatorio_aso,
    renderizar_relatorio_aso_docx,
)
from agente_medico.superficie.web_matriz import pagina_matriz

_PROTOCOLO = carregar(Path(__file__).parent.parent / "protocolo")
_EXAMES = _PROTOCOLO.vocabulario.exames
_CAB = CabecalhoDocumento("CMO", "AURORA", "Adendo", "2026-10-07", "Dra. X", "CRM")


def _sugestao(risco: str, veredito: Any, *, reconhecido: bool = True, exames: tuple[str, ...] = ()) -> SugestaoASO:
    return SugestaoASO(risco, reconhecido, veredito, (Criterio(veredito, "R-ASO-06", "motivo"),), exames)


def _matriz(relatorio: RelatorioASO) -> MatrizGHE:
    return MatrizGHE(
        ghe_id="GHE-01",
        nome_ghe="ALVENARIA",
        linhas=[ExameEmitido(exame="exame_clinico", periodicidade_meses=12)],
        cargos=("Pedreiro", "Servente"),
        sugestao_aso=relatorio,
    )


def _rodapes(caminho: Path) -> list[str]:
    return [p.text for s in Document(str(caminho)).sections for p in s.footer.paragraphs if p.text]


def test_riscos_ordenados_consta_conferir_nao_consta(tmp_path: Path) -> None:
    # Reversão: tirar o `sorted(..., key=_ORDEM)` de montar_relatorio_aso — os
    # riscos saem na ordem do PGR (não consta primeiro).
    relatorio = RelatorioASO((
        _sugestao("umidade", "NAO_CONSTA"),
        _sugestao("Poeira de gesso", "CONFERIR", reconhecido=False),
        _sugestao("ruido", "CONSTA", exames=("audiometria",)),
    ))
    (bloco,) = montar_relatorio_aso([_matriz(relatorio)], _EXAMES).blocos
    assert [ln.sugestao for ln in bloco.linhas] == ["Consta no ASO", "Conferir", "Não consta"]
    assert bloco.linhas[1].risco == "Poeira de gesso (termo do PGR)"
    assert bloco.cargos == ("Pedreiro", "Servente")


def test_documento_traz_tabela_inexistencia_e_aptidao(tmp_path: Path) -> None:
    # Reversões: remover o `if bloco.inexistencia` do render (a frase some);
    # remover o laço de `aptidoes` (a aptidão some); remover a `_tabela` (sem riscos).
    relatorio = RelatorioASO(
        (_sugestao("umidade", "NAO_CONSTA"),),
        aptidoes=("Consignar aptidão para trabalho em altura (NR-35)",),
        inexistencia=True,
    )
    destino = tmp_path / "riscos_aso.docx"
    renderizar_relatorio_aso_docx(montar_relatorio_aso([_matriz(relatorio)], _EXAMES), _CAB, destino)

    documento = Document(str(destino))
    paragrafos = [p.text for p in documento.paragraphs]
    celulas = [c.text for t in documento.tables for r in t.rows for c in r.cells]
    assert TEXTO_INEXISTENCIA in paragrafos
    assert "Consignar aptidão para trabalho em altura (NR-35)" in paragrafos
    assert "umidade" in celulas and "Não consta" in celulas


def test_rodape_confidencial_nos_anexos(tmp_path: Path) -> None:
    # Reversões: remover `aplicar_rodape_confidencial` de renderizar_memorial_docx
    # (memorial sem rodapé) ou de renderizar_relatorio_aso_docx (relatório sem rodapé).
    matriz = _matriz(RelatorioASO((_sugestao("ruido", "CONSTA"),)))
    memorial = tmp_path / "memorial.docx"
    renderizar_memorial_docx(montar_memorial([matriz], _EXAMES, {}), _CAB, memorial)
    aso = tmp_path / "riscos_aso.docx"
    renderizar_relatorio_aso_docx(montar_relatorio_aso([matriz], _EXAMES), _CAB, aso)

    assert _rodapes(memorial) == [RODAPE_CONFIDENCIAL]
    assert _rodapes(aso) == [RODAPE_CONFIDENCIAL]


def test_matriz_assinada_nao_leva_rodape_confidencial(tmp_path: Path) -> None:
    # Reversão: chamar `aplicar_rodape_confidencial` em renderizar_docx — a matriz
    # que vai ao cliente sairia marcada como uso interno.
    matriz = _matriz(RelatorioASO())
    destino = tmp_path / "matriz.docx"
    renderizar_docx(montar_documento([matriz], _EXAMES, _CAB, RodapeDocumento("", "", "")), destino)

    assert RODAPE_CONFIDENCIAL not in _rodapes(destino)


def test_tela_oferece_riscos_para_o_aso(monkeypatch: pytest.MonkeyPatch) -> None:
    # Reversão: remover o `st.download_button` de "riscos_aso.docx" em pagina_matriz
    # (ou não montar `aso_bytes`) — o arquivo não é oferecido.
    from agente_medico.tests.test_web_matriz import _mockar_parse_deterministico, _submeter_formulario

    arquivos: dict[str, bytes] = {}

    def _download_espiao(rotulo: str, dados: Any, **kwargs: Any) -> bool:
        arquivos[str(kwargs.get("file_name"))] = dados
        return False

    monkeypatch.setattr("streamlit.download_button", _download_espiao)
    matriz = _matriz(RelatorioASO((_sugestao("ruido", "CONSTA", exames=("audiometria",)),)))
    _mockar_parse_deterministico(monkeypatch, Resultado(status="OK", matrizes=[matriz]))

    at = AppTest.from_function(pagina_matriz)
    at.run()
    _submeter_formulario(at)

    assert not at.exception
    documento = Document(io.BytesIO(arquivos["riscos_aso.docx"]))
    celulas = [c.text for t in documento.tables for r in t.rows for c in r.cells]
    assert "ruido" in celulas and "Consta no ASO" in celulas
