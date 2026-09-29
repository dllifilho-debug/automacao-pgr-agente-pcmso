"""D-ARQ-87 fatia 2 — memorial de raciocínio da matriz em linguagem clínica: resumo
clínico por regra (nunca o fundamento de auditoria), certeza pelo elo mais fraco que
decide a linha, decisões a confirmar agrupadas por regra, o que cada regra pediu quando
há mais de uma, exames não pedidos, e o download na tela. Cada teste nomeia a reversão
que o deixa vermelho."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from agente_medico.motor.protocolo import carregar
from agente_medico.motor.tipos import ExameEmitido, MatrizGHE, Momento, Motivo, Observacao, Resultado
from agente_medico.superficie.documento_matriz import CabecalhoDocumento
from agente_medico.superficie.memorial_matriz import (
    ROTULO_CERTEZA,
    data_exibicao,
    titulo_ghe,
    linhas_da_tabela,
    montar_memorial,
    nivel_de_certeza,
    renderizar_memorial_docx,
)
from agente_medico.superficie.web_matriz import pagina_matriz

_PROTOCOLO = carregar(Path(__file__).parent.parent / "protocolo")
_EXAMES = _PROTOCOLO.vocabulario.exames
_ADM_PER_MR = frozenset({Momento.ADM, Momento.PER, Momento.MR})
_COM_DEM = _ADM_PER_MR | {Momento.DEM}
_RESUMOS = {r: f"resumo de {r}" for r in ("R-VAL", "R-INT", "R-PSY", "R-ESP-DER", "R-ESP-INT", "R-CLI-01", "R-PKG-ASF")}


def _motivo(regra: str, status: str, meses: int, momentos: frozenset[Momento]) -> Motivo:
    return Motivo(
        regra_id=regra,
        predicado="p",
        risco_origem=None,
        detalhe=None,
        status_regra=status,
        periodicidade_meses=meses,
        momentos=momentos,
        base_normativa=f"fundamento de {regra}",
    )


def _exame(slug: str, *motivos: Motivo) -> ExameEmitido:
    momentos: set[Momento] = set()
    for m in motivos:
        momentos |= m.momentos
    return ExameEmitido(
        exame=slug,
        periodicidade_meses=min(m.periodicidade_meses or 0 for m in motivos),
        momentos=momentos,
        motivos=list(motivos),
    )


def test_toda_regra_ativa_tem_resumo_clinico() -> None:
    # Reversão que mata: apagar o `resumo_clinico` de qualquer regra ativa em regras.yaml
    # (o memorial cairia em "Regra sem resumo clínico").
    sem = [r["id"] for r in _PROTOCOLO.regras if not str(r.get("resumo_clinico", "")).strip()]
    assert sem == []


def test_certeza_e_o_elo_mais_fraco_que_decide_a_linha() -> None:
    # Reversões que matam: certeza = pior status entre TODOS os motivos (o primeiro
    # caso vira interpretação); certeza = melhor status entre todos (o segundo vira validado).
    repete = _exame(
        "espirometria",
        _motivo("R-VAL", "VALIDADO", 24, _COM_DEM),
        _motivo("R-INT", "INTERPRETADO", 24, _COM_DEM),
    )
    unico_dem = _exame(
        "acuidade_visual",
        _motivo("R-VAL", "VALIDADO", 12, _ADM_PER_MR),
        _motivo("R-INT", "INTERPRETADO", 12, _COM_DEM),
    )

    assert ROTULO_CERTEZA[nivel_de_certeza(repete)] == "Protocolo validado pela coordenação"
    assert ROTULO_CERTEZA[nivel_de_certeza(unico_dem)] == "Interpretação do sistema — confirmar"


def test_confirmar_primeiro_agrupa_por_regra_e_ignora_regra_redundante() -> None:
    # Reversões que matam: uma entrada por linha em vez de por regra (duas entradas de
    # R-PSY); `regras_fracas` pegar qualquer motivo INTERPRETADO (R-ESP-INT entraria,
    # embora R-ESP-DER já sustente a linha).
    psy = _motivo("R-PSY", "INTERPRETADO", 12, _ADM_PER_MR)
    esp = (_motivo("R-ESP-DER", "DERIVADO", 24, _COM_DEM), _motivo("R-ESP-INT", "INTERPRETADO", 24, _COM_DEM))
    matrizes = [
        MatrizGHE(ghe_id="GHE-01", linhas=[_exame("avaliacao_psicossocial", psy)]),
        MatrizGHE(ghe_id="GHE-16", linhas=[_exame("avaliacao_psicossocial", psy), _exame("espirometria", *esp)]),
        MatrizGHE(ghe_id="GHE-19", linhas=[_exame("espirometria", *esp)]),
    ]

    memorial = montar_memorial(matrizes, _EXAMES, _RESUMOS)

    (decisao,) = memorial.revisar_primeiro
    assert decisao.regra_id == "R-PSY"
    assert decisao.resumo == "resumo de R-PSY"
    assert decisao.onde == ("Avaliação Psicossocial: GHE-01, GHE-16",)
    assert decisao.codigos == ("GHE-01 · Avaliação Psicossocial", "GHE-16 · Avaliação Psicossocial")


def test_porque_usa_o_resumo_e_mostra_o_que_cada_regra_pediu() -> None:
    # Reversões que matam: usar `base_normativa` no lugar do resumo; tirar o "[pede: ...]"
    # por regra; tirar a frase da menor periodicidade quando as regras divergem.
    clinico = _exame(
        "exame_clinico",
        _motivo("R-CLI-01", "VALIDADO", 12, _ADM_PER_MR),
        _motivo("R-PKG-ASF", "DERIVADO", 6, frozenset({Momento.PER})),
    )
    (bloco,) = montar_memorial([MatrizGHE(ghe_id="GHE-22", linhas=[clinico])], _EXAMES, _RESUMOS).blocos
    (linha,) = bloco.linhas

    assert "resumo de R-CLI-01. [pede: ADM, PER 12 meses, MRO] (ref. R-CLI-01)" in linha.porque
    assert "resumo de R-PKG-ASF. [pede: PER 6 meses] (ref. R-PKG-ASF)" in linha.porque
    assert "vale a menor periodicidade (6 meses)" in linha.porque
    assert "fundamento de" not in linha.porque


def test_tabela_do_ghe_leva_so_a_primeira_frase_e_a_secao_3_o_resumo_inteiro() -> None:
    # Reversões que matam: usar o resumo inteiro em "por que foi pedido" (sem
    # `primeira_frase`); não montar `regras_usadas` (a seção 3 some).
    resumos = {"R-CLI-01": "Todos → clínico anual. NR-07, piso do PCMSO."}
    clinico = _exame("exame_clinico", _motivo("R-CLI-01", "VALIDADO", 12, _ADM_PER_MR))
    memorial = montar_memorial([MatrizGHE(ghe_id="GHE-01", linhas=[clinico])], _EXAMES, resumos)

    (linha,) = memorial.blocos[0].linhas
    assert linha.porque == "Todos → clínico anual. (ref. R-CLI-01)"
    (regra,) = memorial.regras_usadas
    assert regra.resumo == "Todos → clínico anual. NR-07, piso do PCMSO."
    assert regra.certeza == "Protocolo validado pela coordenação"


def test_exames_com_o_mesmo_motivo_viram_uma_linha_na_tabela() -> None:
    # Reversão que mata: uma linha de tabela por exame em `linhas_da_tabela`.
    ativ = _motivo("R-VAL", "VALIDADO", 12, _ADM_PER_MR)
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("hemograma", ativ), _exame("glicemia", ativ), _exame("avaliacao_psicossocial", _motivo("R-PSY", "INTERPRETADO", 12, _ADM_PER_MR))],
    )
    (bloco,) = montar_memorial([matriz], _EXAMES, _RESUMOS).blocos

    tabela = linhas_da_tabela(bloco)
    assert len(tabela) == 2
    assert tabela[0][0].count("\n") == 1


def test_memorial_lista_os_exames_nao_pedidos() -> None:
    # Reversão que mata: não levar `matriz.observacoes` para `nao_pedidos`.
    obs = Observacao(
        regra_id="R-BIO-04-xileno",
        regra_dispensa="R-BIO-05",
        agente="xileno",
        nivel_risco="IRRELEVANTE",
        exames_dispensados=("acido_metilhipurico",),
    )
    matriz = MatrizGHE(
        ghe_id="GHE-18",
        linhas=[_exame("exame_clinico", _motivo("R-CLI-01", "VALIDADO", 12, _ADM_PER_MR))],
        observacoes=(obs,),
    )
    (bloco,) = montar_memorial([matriz], _EXAMES, _RESUMOS).blocos

    (texto,) = bloco.nao_pedidos
    assert texto.startswith("Ácido metil-hipúrico na urina: não pedido — xileno com risco irrelevante")


def test_docx_tem_confirmar_primeiro_e_tabela_por_ghe_sem_fundamento_tecnico(tmp_path: Path) -> None:
    # Reversões que matam: tirar a tabela "confirmar primeiro"; tirar a coluna Correção;
    # tirar a tabela da seção 3; voltar a pôr o fundamento de auditoria
    # (`base_normativa`) no documento.
    psy = _motivo("R-PSY", "INTERPRETADO", 12, _ADM_PER_MR)
    matrizes = [MatrizGHE(ghe_id="GHE-01", linhas=[_exame("avaliacao_psicossocial", psy)])]
    destino = tmp_path / "memorial.docx"
    cab = CabecalhoDocumento("E", "O", "Atualização", "27/09/2026", "Dra. X", "CRM")

    renderizar_memorial_docx(montar_memorial(matrizes, _EXAMES, _RESUMOS), cab, destino)

    documento = Document(str(destino))
    confirmar, ghe, regras = documento.tables
    assert confirmar.rows[1].cells[0].text == "resumo de R-PSY (ref. R-PSY)"
    assert regras.rows[1].cells[0].text == "resumo de R-PSY (ref. R-PSY)"
    assert ghe.rows[0].cells[-1].text == "Correção"
    assert ghe.rows[1].cells[-1].text == ""
    textos = [p.text for p in documento.paragraphs] + [c.text for t in documento.tables for r in t.rows for c in r.cells]
    assert not any("fundamento de" in t for t in textos)


def test_regra_em_todos_os_ghes_aparece_como_todos() -> None:
    # Reversões que matam: `_lista_ghes` sempre listar os GHEs (a primeira asserção
    # vira a lista de 3); tirar a guarda `len(todos) > 1` (o memorial de um GHE só
    # diria "todos os GHEs (1)"). O caso parcial está no teste de "confirmar primeiro".
    psy = _motivo("R-PSY", "INTERPRETADO", 12, _ADM_PER_MR)
    tres = [MatrizGHE(ghe_id=g, linhas=[_exame("avaliacao_psicossocial", psy)]) for g in ("GHE-01", "GHE-02", "GHE-03")]

    (decisao,) = montar_memorial(tres, _EXAMES, _RESUMOS).revisar_primeiro
    (sozinho,) = montar_memorial(tres[:1], _EXAMES, _RESUMOS).revisar_primeiro

    assert decisao.onde == ("Avaliação Psicossocial: todos os GHEs (3)",)
    assert sozinho.onde == ("Avaliação Psicossocial: GHE-01",)


@pytest.mark.parametrize(
    "ghe_id,nome,esperado",
    [
        ("GHE-01", "GHE 01 - ADMINISTRAÇÃO", "GHE-01 — ADMINISTRAÇÃO"),
        ("GHE-16", "GHE 16 – SERRALHERIA", "GHE-16 — SERRALHERIA"),
        ("GHE-07", "SUPERVISÃO DE ATIVIDADES EM OBRA", "GHE-07 — SUPERVISÃO DE ATIVIDADES EM OBRA"),
        ("GHE-01", "GHE 010 - OUTRO", "GHE-01 — GHE 010 - OUTRO"),
        ("3", "ARMAÇÃO", "GHE 3 — ARMAÇÃO"),
    ],
)
def test_titulo_do_ghe_nao_repete_o_codigo(ghe_id: str, nome: str, esperado: str) -> None:
    # Reversões que matam: não tirar o código do nome (os dois primeiros repetem
    # "GHE 01"); tirar o `(?!\d)` (o GHE-01 come o início de "GHE 010"); tirar o
    # prefixo "GHE " de id sem ele (o último vira "3 — ARMAÇÃO").
    assert titulo_ghe(ghe_id, nome) == esperado


def test_data_iso_vira_dia_mes_ano_e_o_resto_sai_como_digitado() -> None:
    # Reversão que mata: `data_exibicao` devolver o texto sem converter (a primeira) ou
    # converter qualquer coisa com hífen (a terceira).
    assert data_exibicao("2026-09-26") == "26/09/2026"
    assert data_exibicao("27/09/2026") == "27/09/2026"
    assert data_exibicao("set-2026") == "set-2026"


def test_docx_tem_titulo_do_ghe_data_e_legenda_dos_momentos(tmp_path: Path) -> None:
    # Reversões que matam: o render voltar a montar o título como `id + nome`; não
    # passar a data por `data_exibicao`; tirar o parágrafo da legenda dos momentos.
    psy = _motivo("R-PSY", "INTERPRETADO", 12, _ADM_PER_MR)
    matrizes = [MatrizGHE(ghe_id="GHE-01", nome_ghe="GHE 01 - ADMINISTRAÇÃO", linhas=[_exame("avaliacao_psicossocial", psy)])]
    destino = tmp_path / "memorial.docx"
    cab = CabecalhoDocumento("CMO", "AURORA", "Adendo", "2026-09-26", "Dra. X", "CRM")

    renderizar_memorial_docx(montar_memorial(matrizes, _EXAMES, _RESUMOS), cab, destino)

    paragrafos = [p.text for p in Document(str(destino)).paragraphs]
    assert "GHE-01 — ADMINISTRAÇÃO" in paragrafos
    assert "Empresa: CMO · Obra: AURORA · Data: 26/09/2026" in paragrafos
    assert any(p.startswith("Momentos: ADM admissional;") and "MRO mudança de riscos ocupacionais" in p for p in paragrafos)


def test_tela_oferece_o_memorial_com_os_resumos_do_protocolo(monkeypatch: pytest.MonkeyPatch) -> None:
    # Reversões que matam: remover o `st.download_button` do memorial; passar resumos
    # vazios a `montar_memorial` em pagina_matriz (sai "Regra sem resumo clínico").
    from agente_medico.tests.test_web_matriz import _mockar_parse_deterministico, _submeter_formulario

    arquivos: dict[str, bytes] = {}

    def _download_espiao(rotulo: str, dados: Any, **kwargs: Any) -> bool:
        arquivos[str(kwargs.get("file_name"))] = dados
        return False

    monkeypatch.setattr("streamlit.download_button", _download_espiao)
    exame = _exame("exame_clinico", _motivo("R-CLI-01", "VALIDADO", 12, _ADM_PER_MR))
    matriz = MatrizGHE(ghe_id="GHE-01", linhas=[exame], cargos=("Cargo Teste",))
    _mockar_parse_deterministico(monkeypatch, Resultado(status="OK", matrizes=[matriz]))

    at = AppTest.from_function(pagina_matriz)
    at.run()
    _submeter_formulario(at)

    assert not at.exception
    documento = Document(io.BytesIO(arquivos["memorial.docx"]))
    celulas = [c.text for t in documento.tables for r in t.rows for c in r.cells]
    resumo = next(str(r["resumo_clinico"]) for r in _PROTOCOLO.regras if r["id"] == "R-CLI-01")
    assert any(resumo in c for c in celulas)
