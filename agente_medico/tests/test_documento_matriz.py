from __future__ import annotations

import time
from datetime import date
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from agente_medico.adaptadores.orquestracao_pgr import processar_arquivo_pgr
from agente_medico.motor.protocolo import carregar
from agente_medico.motor.tipos import EnvelopeConfirmado, ExameEmitido, GHEVerbatim, MatrizGHE, Momento
from agente_medico.superficie.documento_matriz import (
    _COR_DESTAQUE,
    _COR_DESTAQUE_HEX,
    _COR_TEXTO_SOBRE_DESTAQUE,
    _ROTULO_MOMENTO,
    CabecalhoDocumento,
    RodapeDocumento,
    montar_documento,
    renderizar_docx,
    renderizar_html,
)


def _cabecalho() -> CabecalhoDocumento:
    return CabecalhoDocumento(
        empresa="Empresa Teste",
        obra="Obra Teste",
        tipo_documento="Atualização",
        data="01/08/2026",
        medico_coordenador="Dra. Teste",
        crm="CRM-GO 0000",
    )


def _rodape() -> RodapeDocumento:
    return RodapeDocumento(
        responsavel_preenchimento="Teste",
        medico_validador="Dra. Teste",
        data_pgr="01/01/2026",
    )


def _exame(
    slug: str, periodicidade: int = 12, momentos: set[Momento] | None = None
) -> ExameEmitido:
    return ExameEmitido(
        exame=slug,
        periodicidade_meses=periodicidade,
        momentos=momentos if momentos is not None else {Momento.ADM},
    )


# ---------------------------------------------------------------------------
# 003.EO EMENDA 1 — ordem_exibicao (2 dos 3 testes; a unicidade mora em
# test_vocabulario.py por ser checagem pura de yaml).
# ---------------------------------------------------------------------------


def test_ordem_das_celulas_segue_ordem_exibicao_do_vocabulario() -> None:
    # Reversão que mata: ordenar por `exame.exame` (slug alfabético) em vez de
    # por `ordem_exibicao`. "zebra" (ordem 1) vem alfabeticamente DEPOIS de
    # "abacate" (ordem 2) — só a ordem_exibicao acerta a posição.
    vocab = {
        "zebra": {"nome_exibicao": "Zebra", "ordem_exibicao": 1},
        "abacate": {"nome_exibicao": "Abacate", "ordem_exibicao": 2},
    }
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("abacate"), _exame("zebra")],
        cargos=("Cargo Único",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celulas = doc.blocos[0].linhas[0].celulas
    assert celulas[0].startswith("Zebra")
    assert celulas[1].startswith("Abacate")


def test_exame_sem_ordem_exibicao_vai_para_o_fim() -> None:
    # Reversão que mata: fallback devolver 0 em vez de sentinela-por-tag alta
    # -> "aaa_sem_ordem" (sem campo) subiria para o topo por ser 0 == menor
    # ordem_exibicao possível. Aqui ele tem que ficar depois de "so_com_ordem"
    # (ordem_exibicao=5) mesmo perdendo no alfabeto.
    vocab = {
        "so_com_ordem": {"nome_exibicao": "Com Ordem", "ordem_exibicao": 5},
        "aaa_sem_ordem": {"nome_exibicao": "Sem Ordem"},
    }
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("aaa_sem_ordem"), _exame("so_com_ordem")],
        cargos=("Cargo Único",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celulas = doc.blocos[0].linhas[0].celulas
    assert celulas[0].startswith("Com Ordem")
    assert celulas[1].startswith("Sem Ordem")


# ---------------------------------------------------------------------------
# Fatia 2 do prompt original — expansão GHE→cargo + emissor HTML.
# ---------------------------------------------------------------------------


def test_expansao_ghe_cargo_replica_celulas() -> None:
    # Reversão que mata: só o primeiro cargo recebe as células (ex.: `linhas =
    # (LinhaCargo(matriz.cargos[0], celulas),)` em vez de um por cargo).
    vocab = {
        "exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1},
        "audiometria": {"nome_exibicao": "Audiometria", "ordem_exibicao": 2},
        "hemograma": {"nome_exibicao": "Hemograma", "ordem_exibicao": 3},
    }
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("exame_clinico"), _exame("audiometria"), _exame("hemograma")],
        cargos=("Carpinteiro", "Ajudante", "Encarregado", "Estagiário"),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    bloco = doc.blocos[0]
    assert len(bloco.linhas) == 4
    celulas_esperadas = bloco.linhas[0].celulas
    assert len(celulas_esperadas) == 3
    for linha in bloco.linhas:
        assert linha.celulas == celulas_esperadas


def test_ghe_sem_cargos_nao_inventa_linha() -> None:
    # Reversão que mata: inserir um placeholder "(sem cargo)" no ramo vazio.
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=())
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    assert doc.blocos[0].linhas == ()


def test_mapa_momentos_cobre_todos_os_membros_do_enum() -> None:
    # Reversão que mata: remover uma entrada do mapa (ex.: RT).
    assert set(_ROTULO_MOMENTO.keys()) == set(Momento)


def test_celula_usa_nome_exibicao_do_vocabulario() -> None:
    # Reversão que mata: usar exame.exame (o slug) em vez do nome_exibicao.
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(
        ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=("Cargo",)
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert "Exame Clínico" in celula
    assert "exame_clinico" not in celula


# DT-003EW-02 — periodicidade impressa na celula. Regra de 003.EO, medida
# contra o acervo em 29/08/2026 (28 documentos, instrumento
# scripts/medir_cobertura_e_forma.py): 2896 confirmam / 122 contrariam nos
# 10 documentos de 2026, contra 1323/1681 nos de 2025. E a convencao
# corrente do escritorio, nao invariante do acervo historico — o app emite
# documento novo, logo emite a convencao corrente.
# Medicao, corte por documento e ressalvas em
# docs/referencia/MEDICAO_003FE_regra_forma_periodicidade.md.


def test_periodicidade_diferente_de_12_meses_imprime_o_numero() -> None:
    # Reversão que mata: `mostrar` fixo em False (nunca imprime o número).
    vocab = {"espirometria": {"nome_exibicao": "Espirometria"}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[
            _exame(
                "espirometria",
                periodicidade=24,
                momentos={Momento.ADM, Momento.PER, Momento.MR, Momento.DEM},
            )
        ],
        cargos=("Cargo",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert celula == "Espirometria (ADM, PER 24 meses, MRO, DEM)"


def test_periodicidade_de_12_meses_nao_imprime_o_numero() -> None:
    # Reversão que mata: `mostrar` fixo em True (sempre imprime o número).
    vocab = {"audiometria": {"nome_exibicao": "Audiometria"}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[
            _exame(
                "audiometria",
                periodicidade=12,
                momentos={Momento.ADM, Momento.PER, Momento.MR, Momento.DEM},
            )
        ],
        cargos=("Cargo",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert celula == "Audiometria (ADM, PER, MRO, DEM)"


def test_rx_torax_oit_imprime_periodicidade_mesmo_a_12_meses() -> None:
    # Reversão que mata: remover o termo `or bool(entrada.get(
    # "periodicidade_sempre_visivel", False))` de `_formatar_celula` — sem ele a
    # exceção não existe e rx_torax_oit a 12M sai sem número. (Apagar o campo do
    # YAML mata o teste 6, não este: aqui o vocabulário é literal.)
    vocab = {"rx_torax_oit": {"nome_exibicao": "RX Tórax OIT", "periodicidade_sempre_visivel": True}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[
            _exame(
                "rx_torax_oit",
                periodicidade=12,
                momentos={Momento.ADM, Momento.PER, Momento.MR, Momento.DEM},
            )
        ],
        cargos=("Cargo",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert celula == "RX Tórax OIT (ADM, PER 12 meses, MRO, DEM)"


def test_exame_clinico_semestral_imprime_seis_meses() -> None:
    # Reversão que mata: trocar `!= 12` por `> 12` — 6 é menor que 12, ficaria
    # mudo justo no caso que a regra existe para cobrir.
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico"}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[
            _exame(
                "exame_clinico",
                periodicidade=6,
                momentos={Momento.ADM, Momento.PER, Momento.MR, Momento.RT, Momento.DEM},
            )
        ],
        cargos=("Cargo",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert celula == "Exame Clínico (ADM, PER 6 meses, MRO, RET, DEM)"


def test_numero_so_gruda_no_momento_per() -> None:
    # Reversão que mata: grudar o número no primeiro momento presente em vez
    # de especificamente no PER — aqui não há PER, então nenhum número sai.
    vocab = {"rx_coluna_lombo_sacra": {"nome_exibicao": "RX Coluna Lombo-Sacra"}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[
            _exame(
                "rx_coluna_lombo_sacra",
                periodicidade=24,
                momentos={Momento.ADM, Momento.MR},
            )
        ],
        cargos=("Cargo",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    celula = doc.blocos[0].linhas[0].celulas[0]
    assert celula == "RX Coluna Lombo-Sacra (ADM, MRO)"


def test_periodicidade_sempre_visivel_e_exatamente_rx_torax_oit() -> None:
    # Reversão que mata: acrescentar o campo em outro exame do vocabulário
    # real sem medir a exceção — computado sobre o YAML de produção (D-ARQ-67).
    from agente_medico.motor.protocolo import carregar

    protocolo = carregar(Path(__file__).parent.parent / "protocolo")
    marcados = {
        slug
        for slug, entrada in protocolo.vocabulario.exames.items()
        if entrada.get("periodicidade_sempre_visivel")
    }
    assert marcados == {"rx_torax_oit"}


def test_html_escapa_conteudo_de_dado() -> None:
    # Reversão que mata: remover a chamada de escape (montar a célula/cargo
    # direto na f-string do HTML).
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("exame_clinico")],
        cargos=("<script>alert(1)</script>",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())
    saida = renderizar_html(doc)
    assert "<script>alert(1)</script>" not in saida
    assert "&lt;script&gt;" in saida


# ---------------------------------------------------------------------------
# Fatia 3 — emissor DOCX. Mesma DocumentoMatriz da fatia 2, não recalcula nada.
# ---------------------------------------------------------------------------


def test_docx_tem_uma_tabela_por_ghe(tmp_path: Path) -> None:
    # Reversão que mata: emitir uma tabela única concatenando todos os GHEs
    # (ex.: `add_table` só uma vez, fora do loop de blocos).
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    m1 = MatrizGHE(ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=("A",))
    m2 = MatrizGHE(ghe_id="GHE-02", linhas=[_exame("exame_clinico")], cargos=("B",))
    m3 = MatrizGHE(ghe_id="GHE-03", linhas=[_exame("exame_clinico")], cargos=("C",))
    doc = montar_documento([m1, m2, m3], vocab, _cabecalho(), _rodape())

    destino = tmp_path / "matriz.docx"
    renderizar_docx(doc, destino)

    reaberto = DocxDocument(str(destino))
    assert len(reaberto.tables) == 3


def test_docx_celula_de_exames_preserva_quebra_por_exame(tmp_path: Path) -> None:
    # Reversão que mata: juntar os exames com ", ".join(...) num parágrafo só
    # em vez de um add_paragraph() por exame -> a célula teria 1 parágrafo
    # para N exames, não N. Este é o teste que discrimina de verdade: um
    # emissor ingênuo passaria em "uma tabela por GHE" e erraria aqui.
    vocab = {
        "exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1},
        "audiometria": {"nome_exibicao": "Audiometria", "ordem_exibicao": 2},
        "hemograma": {"nome_exibicao": "Hemograma", "ordem_exibicao": 3},
    }
    matriz = MatrizGHE(
        ghe_id="GHE-01",
        linhas=[_exame("exame_clinico"), _exame("audiometria"), _exame("hemograma")],
        cargos=("Carpinteiro",),
    )
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())

    destino = tmp_path / "matriz.docx"
    renderizar_docx(doc, destino)

    reaberto = DocxDocument(str(destino))
    tabela = reaberto.tables[0]
    celula_exames = tabela.rows[1].cells[1]
    assert len(celula_exames.paragraphs) == 3
    textos = [p.text for p in celula_exames.paragraphs]
    assert textos[0].startswith("Exame Clínico")
    assert textos[1].startswith("Audiometria")
    assert textos[2].startswith("Hemograma")


# ---------------------------------------------------------------------------
# D-ARQ-73, nota de aplicação desta sessão — estilo visual portado de
# modulo_pcmso.py::gerar_docx_rq61 (legado). Conteúdo/contagem de tabelas e
# linhas não muda; só a aparência.
# ---------------------------------------------------------------------------


def test_docx_tabela_por_ghe_usa_estilo_com_borda(tmp_path: Path) -> None:
    # Reversão que mata: tirar `tabela.style = "Table Grid"` -> volta ao
    # estilo default do python-docx (sem nome, sem borda visível).
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=("A",))
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())

    destino = tmp_path / "matriz.docx"
    renderizar_docx(doc, destino)

    reaberto = DocxDocument(str(destino))
    tabela = reaberto.tables[0]
    assert tabela.style is not None
    assert tabela.style.name == "Table Grid"


def test_docx_cabecalho_de_coluna_tem_fundo_e_texto_destacados(tmp_path: Path) -> None:
    # Reversão que mata: tirar a chamada de `_aplicar_fundo` e a cor do texto
    # do cabeçalho de coluna -> célula fica sem `w:shd` no XML e o texto sai
    # na cor padrão (preto), não branco.
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=("A",))
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())

    destino = tmp_path / "matriz.docx"
    renderizar_docx(doc, destino)

    reaberto = DocxDocument(str(destino))
    tabela = reaberto.tables[0]
    celula_funcao = tabela.rows[0].cells[0]

    tcPr = celula_funcao._tc.tcPr
    assert tcPr is not None
    sombreado = tcPr.find(qn("w:shd"))
    assert sombreado is not None
    assert sombreado.get(qn("w:fill")).upper() == _COR_DESTAQUE_HEX

    run_cabecalho = celula_funcao.paragraphs[0].runs[0]
    assert run_cabecalho.bold is True
    assert run_cabecalho.font.color.rgb == _COR_TEXTO_SOBRE_DESTAQUE


def test_docx_titulo_e_cabecalho_de_ghe_usam_cor_de_destaque(tmp_path: Path) -> None:
    # Reversão que mata: tirar `run_titulo.font.color.rgb = _COR_DESTAQUE` (e
    # o equivalente no cabeçalho de GHE) -> cor volta a None (preto padrão).
    vocab = {"exame_clinico": {"nome_exibicao": "Exame Clínico", "ordem_exibicao": 1}}
    matriz = MatrizGHE(ghe_id="GHE-01", linhas=[_exame("exame_clinico")], cargos=("A",))
    doc = montar_documento([matriz], vocab, _cabecalho(), _rodape())

    destino = tmp_path / "matriz.docx"
    renderizar_docx(doc, destino)

    reaberto = DocxDocument(str(destino))
    titulo_documento = reaberto.paragraphs[0]
    assert titulo_documento.runs[0].font.color.rgb == _COR_DESTAQUE

    cabecalhos_ghe = [p for p in reaberto.paragraphs if p.style.name == "Heading 2"]
    assert len(cabecalhos_ghe) == 1
    assert cabecalhos_ghe[0].runs[0].font.color.rgb == _COR_DESTAQUE


# ---------------------------------------------------------------------------
# 003.EP fatia 3 — e2e real: PDF Fascino -> parse determinístico -> hidratação
# -> processar_arquivo_pgr -> montar_documento. Fecha o critério de pronto do
# S2 (D-ARQ-65 fatia 1/2, 003.EP): os 41 cargos (não mais 19 células
# combinadas) têm de sobreviver até a estrutura que o emissor HTML/DOCX
# consome — as fatias 1-2 mediram só até GHEVerbatim, nunca até aqui.
# ---------------------------------------------------------------------------

_MATRIZES_DIR = Path(__file__).parent.parent.parent / "matrizes_originais"
_PDF_FASCINO = (
    _MATRIZES_DIR / "PGR - CONSCIENTE CONSTRUTORA E INCORPORADORA SPE 0030 - FASCINO  (15.07.26).pdf"
)
_PROTOCOLO_DIR = Path(__file__).parent.parent / "protocolo"

requer_pdfs = pytest.mark.skipif(
    not _PDF_FASCINO.exists(),
    reason="PDF Fascino ausente; harness integração e2e (003.EP fatia 3) indisponível",
)


class _TranscritorGHENuncaChamado:
    """Cliente-bomba (molde MockTranscritorGHENuncaChamado, test_orquestracao_pgr.py):
    a rota determinística (D-ARQ-65) tem que ser aceita para o Fascino sem
    invocar LLM — falha alto e explícito em vez de mascarar em silêncio uma
    invocação indevida."""

    def transcrever(self, bloco: str) -> GHEVerbatim:
        raise AssertionError("cliente LLM GHE não deveria ser invocado — rota determinística aceita")


class _TranscritorCardNuncaChamado:
    """Dummy para a rota ghe (molde MockTranscritorCardNuncaChamado): se
    processar_arquivo_pgr chamar cliente_card fora da rota card, é bug de
    roteamento."""

    def transcrever(self, card: str, titulo: str) -> GHEVerbatim:
        raise AssertionError("cliente_card não deveria ser invocado na rota ghe")


@requer_pdfs
def test_pipeline_real_fascino_ate_documento_41_linhas_cargo() -> None:
    protocolo = carregar(_PROTOCOLO_DIR)
    envelope = EnvelopeConfirmado(validade=date.today(), assinatura_engenheiro=True)

    inicio = time.monotonic()
    resultado, _pendencias = processar_arquivo_pgr(
        _PDF_FASCINO,
        protocolo,
        _TranscritorGHENuncaChamado(),
        _TranscritorCardNuncaChamado(),
        envelope=envelope,
    )
    assert resultado is not None

    doc = montar_documento(
        resultado.matrizes, protocolo.vocabulario.exames, _cabecalho(), _rodape()
    )
    duracao = time.monotonic() - inicio
    if duracao > 150:
        print(f"AVISO 003.EP fatia 3: pipeline e2e Fascino levou {duracao:.1f}s (> 150s)")

    total_linhas_cargo = sum(len(bloco.linhas) for bloco in doc.blocos)
    assert total_linhas_cargo == 41

    cargos_extraidos = {linha.cargo for bloco in doc.blocos for linha in bloco.linhas}
    for cargo_recuperado in (
        "Encarregado de Pintor",
        "Encarregado de Carpinteiro",
        "Supervisor de Instalações Elétricas",
        "Auxiliar de Obra",
        "Aprendiz Administrativo de Obra",
        "Assistente Administrativo de Obras",
    ):
        assert cargo_recuperado in cargos_extraidos

    assert all(linha.cargo != "" for bloco in doc.blocos for linha in bloco.linhas)
