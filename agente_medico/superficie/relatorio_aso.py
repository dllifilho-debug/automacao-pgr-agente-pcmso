"""Relatório "riscos para o ASO" (D-ARQ-91 fatia 2): anexo para as médicas, ao lado da
matriz e do memorial. Por GHE, com os cargos: cada risco com a sugestão (consta no ASO,
conferir, não consta), o motivo com a regra R-ASO e os exames da matriz que o risco
motiva. Apresentação-pura sobre `MatrizGHE.sugestao_aso` (D-ARQ-72): nenhuma decisão aqui.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from agente_medico.motor.tipos import GHEPGR, PGR, MatrizGHE, SugestaoASO, VereditoASO
from agente_medico.superficie.documento_matriz import (
    _COR_DESTAQUE,
    CabecalhoDocumento,
    _sanitizar,
    nome_ghe_exibicao,
    titulo_ghe,
)
from agente_medico.superficie.memorial_matriz import (
    TITULO_AGRAVOS,
    _agente_exibicao,
    _nome_exame,
    _tabela,
    _termo_exibicao,
    agravos_do_ghe,
    aplicar_rodape_confidencial,
    data_exibicao,
)

ROTULO_VEREDITO: dict[VereditoASO, str] = {
    "CONSTA": "Consta no ASO",
    "CONFERIR": "Conferir",
    "NAO_CONSTA": "Não consta",
}
_ORDEM: dict[VereditoASO, int] = {"CONSTA": 0, "CONFERIR": 1, "NAO_CONSTA": 2}

TEXTO_INEXISTENCIA = (
    "Sugestão para o ASO: inexistência de perigos ou fatores de risco que necessitem de "
    "controle médico previsto no PCMSO (NR-07 item 7.5.19.1 \"c\")."
)

_COLUNAS: tuple[tuple[str, float], ...] = (
    ("Risco", 4.0),
    ("Sugestão", 2.4),
    ("Motivo", 6.6),
    ("Agravo à saúde (PGR)", 6.6),
    ("Exames na matriz", 2.8),
    ("Correção", 2.5),
)


@dataclass(frozen=True)
class LinhaRisco:
    risco: str
    sugestao: str
    motivo: str
    exames: str
    agravo: str = "—"


@dataclass(frozen=True)
class BlocoRiscos:
    ghe_id: str
    nome_ghe: str
    cargos: tuple[str, ...]
    linhas: tuple[LinhaRisco, ...]
    aptidoes: tuple[str, ...]
    inexistencia: bool
    agravos: tuple[str, ...] = ()


@dataclass(frozen=True)
class RelatorioRiscosASO:
    blocos: tuple[BlocoRiscos, ...]
    contagem: tuple[tuple[str, int], ...]


def _agravo_pgr(s: SugestaoASO, ghe: GHEPGR | None) -> str:
    """D-ARQ-93: agravo que o PGR escreve para o risco — casado pelo agente resolvido ou,
    no termo não reconhecido, pelo termo do PGR. Linhas do PGR do mesmo agente com agravos
    diferentes saem todas; nenhum agravo lido: "—" (D-ARQ-22)."""
    if ghe is None:
        return "—"
    agravos = (
        _termo_exibicao(r.agravo)
        for r in ghe.riscos
        if r.agravo and ((r.agente == s.risco) if s.reconhecido else (r.agente is None and r.termo == s.risco))
    )
    return "\n".join(dict.fromkeys(agravos)) or "—"


def _linha(s: SugestaoASO, exames_vocab: dict[str, Any], ghe: GHEPGR | None = None) -> LinhaRisco:
    risco = _agente_exibicao(s.risco) if s.reconhecido else f"{_termo_exibicao(s.risco)} (termo do PGR)"
    return LinhaRisco(
        risco=risco,
        sugestao=ROTULO_VEREDITO[s.veredito],
        motivo=_sanitizar("\n".join(f"{c.texto} (ref. {c.regra})" for c in s.criterios)),
        exames=", ".join(_nome_exame(e, exames_vocab) for e in s.exames) or "—",
        agravo=_agravo_pgr(s, ghe),
    )


def montar_relatorio_aso(
    matrizes: Sequence[MatrizGHE], exames_vocab: dict[str, Any], pgr: PGR | None = None
) -> RelatorioRiscosASO:
    ghes_pgr = {g.id: g for g in pgr.ghes} if pgr is not None else {}
    contagem: Counter[VereditoASO] = Counter()
    blocos: list[BlocoRiscos] = []
    for matriz in matrizes:
        sugestoes = sorted(matriz.sugestao_aso.riscos, key=lambda s: _ORDEM[s.veredito])
        contagem.update(s.veredito for s in sugestoes)
        blocos.append(
            BlocoRiscos(
                ghe_id=matriz.ghe_id,
                nome_ghe=nome_ghe_exibicao(matriz.nome_ghe),
                cargos=tuple(_sanitizar(c) for c in matriz.cargos),
                linhas=tuple(_linha(s, exames_vocab, ghes_pgr.get(matriz.ghe_id)) for s in sugestoes),
                aptidoes=matriz.sugestao_aso.aptidoes,
                inexistencia=matriz.sugestao_aso.inexistencia,
                agravos=agravos_do_ghe(ghes_pgr[matriz.ghe_id]) if matriz.ghe_id in ghes_pgr else (),
            )
        )
    return RelatorioRiscosASO(
        blocos=tuple(blocos),
        contagem=tuple((ROTULO_VEREDITO[v], contagem[v]) for v in _ORDEM if contagem[v]),
    )


def renderizar_relatorio_aso_docx(
    relatorio: RelatorioRiscosASO, cabecalho: CabecalhoDocumento, destino: Path
) -> None:
    documento = Document()
    for secao in documento.sections:
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = secao.page_height, secao.page_width
        secao.top_margin = secao.bottom_margin = Cm(1.5)
        secao.left_margin = secao.right_margin = Cm(1.5)
    aplicar_rodape_confidencial(documento)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("Riscos para o ASO — sugestão")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = _COR_DESTAQUE
    documento.add_paragraph(
        f"Empresa: {cabecalho.empresa} · Obra: {cabecalho.obra} · Data: {data_exibicao(cabecalho.data)}"
    )
    documento.add_paragraph(
        "O ASO traz \"a descrição dos perigos ou fatores de risco identificados e classificados no PGR que "
        "necessitem de controle médico previsto no PCMSO, ou a sua inexistência\" (NR-07 item 7.5.19.1 \"c\"). "
        "O PCMSO continua com todos os riscos; esta é a sugestão de quais constam no ASO, pelos critérios da "
        "coordenação e da NR-07, para a médica validar. \"Conferir\": o sistema não tem o dado para decidir. "
        "Os exames aparecem só como informação — no ASO vai o risco. Para corrigir, escreva na coluna Correção."
    )

    documento.add_heading("Resumo", level=1)
    for rotulo, n in relatorio.contagem:
        documento.add_paragraph(f"{rotulo}: {n}", style="List Bullet")

    for bloco in relatorio.blocos:
        cabecalho_ghe = documento.add_heading(titulo_ghe(bloco.ghe_id, bloco.nome_ghe), level=2)
        if cabecalho_ghe.runs:
            cabecalho_ghe.runs[0].font.color.rgb = _COR_DESTAQUE
        documento.add_paragraph(f"Cargos: {', '.join(bloco.cargos) or '—'}")
        if bloco.linhas:
            _tabela(
                documento,
                _COLUNAS,
                [(ln.risco, ln.sugestao, ln.motivo, ln.agravo, ln.exames, "") for ln in bloco.linhas],
            )
        if bloco.inexistencia:
            documento.add_paragraph(TEXTO_INEXISTENCIA)
        for aptidao in bloco.aptidoes:
            documento.add_paragraph(aptidao, style="List Bullet")
        if bloco.agravos:
            documento.add_paragraph(TITULO_AGRAVOS)
            for texto in bloco.agravos:
                documento.add_paragraph(texto, style="List Bullet")
    documento.save(str(destino))
