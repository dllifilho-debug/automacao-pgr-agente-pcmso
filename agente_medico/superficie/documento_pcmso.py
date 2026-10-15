"""PCMSO inteiro gerado na mesma rodada da matriz (D-ARQ-96, etapa 3, fatia 1) — apresentação pura.

Parte do modelo `.docx` do Seconci 2026 (`modelos/pcmso_seconci_2026.docx`, texto fixo das seções 3–16) e preenche
os campos `{{NOME}}` e o bloco de cada GHE: cargos, tabela de perigos com o agravo de `agravos.yaml` e tabela de
exames com o código Seconci de `exames.yaml`. Os exames são os mesmos da matriz por construção — saem de
`MatrizGHE.linhas`, sem decisão clínica nova (D-ARQ-96 cl.5). Dado que o app não tem fica em branco visível e vai
para a lista de avisos, nunca inventado.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.text.paragraph import Paragraph

from agente_medico.motor.identificacao_pgr import cnpj_valido
from agente_medico.motor.tipos import GHEPGR, PGR, ExameEmitido, MatrizGHE, Momento, RiscoPGR
from agente_medico.superficie.agravos_pcmso import OrigemAgravo, VocabularioAgravos, agravo_do_perigo
from agente_medico.superficie.documento_matriz import (
    CabecalhoDocumento,
    _chave_ordem_exame,
    _formatar_observacao,
    titulo_ghe_rq61,
)

MODELO_PADRAO = Path(__file__).resolve().parent / "modelos" / "pcmso_seconci_2026.docx"

EM_BRANCO = "____________"
_MARCA_GHE = "{{BLOCOS_GHE}}"
_CAMPO = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
_LARGURA_TOTAL = 10467
_LARGURAS_PERIGOS = (4186, 2094, 4187)
_LARGURAS_EXAMES = (3488, 1163, 1164, 1163, 1163, 1163, 1163)
_FUNDO_CABECALHO = "D9D9D9"
_COLUNAS_EXAMES = ("Exames", "ADMISSÃO", "APÓS ADM.", "PERIÓDICO", "RET. TRAB", "MUD. FUNÇ.", "DEMISSÃO")
_COLUNAS_PERIGOS = ("Perigo / Fator de Risco", "Grupo", "Descrições dos perigos e possíveis lesões ou agravos à saúde")
_SUFIXO_PGR = " (do PGR — conferir)"
_SEM_AGRAVO = "(sem agravo no PGR — conferir)"
_MESES = ("JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO",
          "NOVEMBRO", "DEZEMBRO")
# Rótulo do grupo como no modelo; a chave é o radical sem acento do texto que o PGR escreve (D-ARQ-95 emenda).
_GRUPOS = (("FISIC", "Físico"), ("QUIMIC", "Químico"), ("BIOLOGIC", "Biológico"), ("ERGONOMIC", "Ergonômicos"),
           ("ACIDENT", "Acidente"), ("PSICOSSOCIA", "Psicossocial"))
_CONSIDERACOES = {"OBRA NOVA": "Elaboração do PCMSO", "ATUALIZAÇÃO": "Atualização do PCMSO",
                  "ADENDO": "Adendo ao PCMSO", "FUNÇÕES INICIAIS": "PCMSO de funções iniciais"}


@dataclass(frozen=True)
class DadosPCMSO:
    """Identificação do documento. Os seis primeiros vêm do cabeçalho da matriz; os demais, quando houver."""

    razao_social: str
    obra: str
    data_documento: str  # dd/mm/aaaa, como no cabeçalho da matriz
    tipo_documento: str
    medico_nome: str
    medico_crm: str
    versao: str = ""
    endereco: str = ""
    complemento: str = ""
    cnpj: str = ""
    cep: str = ""
    cidade: str = ""
    bairro: str = ""
    uf: str = ""
    cnae: str = ""
    grau_risco: str = ""
    descricao_cnae: str = ""
    medico_nit: str = ""
    medico_cpf: str = ""
    medico_especialidade: str = ""
    medico_rqe: str = ""


def dados_do_cabecalho(cabecalho: CabecalhoDocumento) -> DadosPCMSO:
    """O que a tela da matriz já pede; o resto do PCMSO sai em branco, para completar no Word."""
    return DadosPCMSO(
        razao_social=cabecalho.empresa,
        obra=cabecalho.obra,
        data_documento=cabecalho.data,
        tipo_documento=cabecalho.tipo_documento,
        medico_nome=cabecalho.medico_coordenador,
        medico_crm=cabecalho.crm_formatado().removeprefix("CRM").strip(" :-"),
    )


@dataclass
class AvisosPCMSO:
    """O que a médica e o operador precisam completar ou conferir no PCMSO gerado."""

    campos_em_branco: list[str] = field(default_factory=list)
    exames_sem_codigo: list[str] = field(default_factory=list)
    agravos_do_pgr: int = 0
    agravos_ausentes: int = 0
    ghes_sem_pgr: list[str] = field(default_factory=list)
    cnpj_invalido: bool = False


def linhas_de_aviso(avisos: AvisosPCMSO) -> list[str]:
    """Avisos do PCMSO para a tela, em linguagem de quem completa o documento."""
    linhas: list[str] = []
    if avisos.campos_em_branco:
        linhas.append("Completar no Word (em branco no PCMSO): " + ", ".join(avisos.campos_em_branco) + ".")
    if avisos.exames_sem_codigo:
        linhas.append("Exames sem código Seconci conhecido (saem sem código): "
                      + ", ".join(avisos.exames_sem_codigo) + ".")
    if avisos.agravos_do_pgr or avisos.agravos_ausentes:
        linhas.append(
            f"Agravos a conferir: {avisos.agravos_do_pgr} com o texto do PGR e {avisos.agravos_ausentes} sem agravo "
            "no PGR — estão marcados \"conferir\" na tabela de perigos."
        )
    if avisos.ghes_sem_pgr:
        linhas.append("GHE sem os perigos do PGR: " + ", ".join(avisos.ghes_sem_pgr) + ".")
    if avisos.cnpj_invalido:
        linhas.append("CNPJ com dígito verificador que não confere — conferir no PGR e corrigir no Word.")
    return linhas


@dataclass(frozen=True)
class LinhaPerigo:
    perigo: str
    grupo: str
    agravo: str
    origem: OrigemAgravo


def _sanitizar(texto: str) -> str:
    return " ".join(texto.replace("\x00", " ").split())


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)).upper()


def rotulo_grupo(tipo: str) -> str:
    chave = _sem_acento(tipo)
    return next((rotulo for radical, rotulo in _GRUPOS if radical in chave), _sanitizar(tipo))


def _data(texto: str) -> date | None:
    m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", texto)
    if m is None:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def valores_dos_campos(dados: DadosPCMSO, avisos: AvisosPCMSO) -> dict[str, str]:
    """Valor de cada campo do modelo; vazio vira EM_BRANCO e entra nos avisos com o nome legível."""
    brutos: dict[str, tuple[str, str]] = {
        "RAZAO_SOCIAL": ("razão social", dados.razao_social),
        "OBRA": ("obra", dados.obra),
        "DATA_DOCUMENTO": ("data do documento", dados.data_documento),
        "VERSAO": ("versão", dados.versao),
        "ENDERECO": ("endereço", dados.endereco),
        "COMPLEMENTO": ("complemento", dados.complemento),
        "CNPJ": ("CNPJ", dados.cnpj),
        "CEP": ("CEP", dados.cep),
        "CIDADE": ("cidade", dados.cidade),
        "BAIRRO": ("bairro", dados.bairro),
        "UF": ("UF", dados.uf),
        "CNAE": ("CNAE", dados.cnae),
        "GRAU_RISCO": ("grau de risco", dados.grau_risco),
        "DESCRICAO_CNAE": ("descrição do CNAE", dados.descricao_cnae),
        "REVISAO_DATA": ("data da revisão", dados.data_documento),
        "REVISAO_NUMERO": ("número da revisão", dados.versao),
        "REVISAO_CONSIDERACOES": ("considerações da revisão",
                                  _CONSIDERACOES.get(_sanitizar(dados.tipo_documento).upper(), "")),
        "MEDICO_NOME": ("nome do médico responsável", dados.medico_nome),
        "MEDICO_CRM": ("CRM do médico responsável", dados.medico_crm),
        "MEDICO_NIT": ("NIT do médico responsável", dados.medico_nit),
        "MEDICO_CPF": ("CPF do médico responsável", dados.medico_cpf),
        "MEDICO_ESPECIALIDADE": ("especialidade do médico responsável", dados.medico_especialidade),
        "MEDICO_RQE": ("RQE do médico responsável", dados.medico_rqe),
    }
    valores: dict[str, str] = {}
    for nome, (legivel, valor) in brutos.items():
        valor = _sanitizar(valor)
        if not valor:
            avisos.campos_em_branco.append(legivel)
        valores[nome] = valor or EM_BRANCO
    avisos.cnpj_invalido = bool(_sanitizar(dados.cnpj)) and not cnpj_valido(dados.cnpj)
    inicio = _data(dados.data_documento)
    if inicio is None:
        avisos.campos_em_branco.append("período do cronograma")
        valores["CRONOGRAMA_INICIO"] = valores["CRONOGRAMA_FIM"] = EM_BRANCO
        for n in range(1, 14):
            valores[f"CRONOGRAMA_MES_{n:02d}"] = EM_BRANCO
        return valores
    for n in range(13):
        indice = inicio.month - 1 + n
        mes, ano = _MESES[indice % 12], inicio.year + indice // 12
        valores[f"CRONOGRAMA_MES_{n + 1:02d}"] = f"{mes[:3]} {ano}"
        if n == 0:
            valores["CRONOGRAMA_INICIO"] = f"{mes} DE {ano}"
        if n == 12:
            valores["CRONOGRAMA_FIM"] = f"{mes} DE {ano}"
    return valores


def linhas_de_perigo(ghe: GHEPGR, vocabulario: VocabularioAgravos, avisos: AvisosPCMSO) -> list[LinhaPerigo]:
    """Uma linha por perigo do PGR, sem repetir; os fatores psicossociais numa linha só (decisão 5, D-ARQ-96)."""
    linhas: list[LinhaPerigo] = []
    vistos: set[str] = set()
    psicossociais: list[str] = []
    agravo_psicossocial: LinhaPerigo | None = None
    for risco in ghe.riscos:
        perigo = _perigo(risco)
        if not perigo or perigo.casefold() in vistos:
            continue
        vistos.add(perigo.casefold())
        agravo = agravo_do_perigo(perigo, risco.agravo, vocabulario)
        linha = LinhaPerigo(perigo=perigo, grupo=rotulo_grupo(risco.tipo), agravo=agravo.texto, origem=agravo.origem)
        if agravo.chave == "psicossocial":
            psicossociais.append(perigo)
            agravo_psicossocial = agravo_psicossocial or linha
            continue
        linhas.append(linha)
    if agravo_psicossocial is not None:
        linhas.append(LinhaPerigo(
            perigo="Fatores psicossociais: " + "; ".join(psicossociais),
            grupo=agravo_psicossocial.grupo,
            agravo=agravo_psicossocial.agravo,
            origem=agravo_psicossocial.origem,
        ))
    for linha in linhas:
        if linha.origem is OrigemAgravo.PGR_CONFERIR:
            avisos.agravos_do_pgr += 1
        elif linha.origem is OrigemAgravo.AUSENTE:
            avisos.agravos_ausentes += 1
    return linhas


def _perigo(risco: RiscoPGR) -> str:
    if risco.termo:
        return _sanitizar(risco.termo)
    return risco.agente.replace("_", " ") if risco.agente else ""


def texto_do_agravo(linha: LinhaPerigo) -> str:
    if linha.origem is OrigemAgravo.PGR_CONFERIR:
        return linha.agravo + _SUFIXO_PGR
    if linha.origem is OrigemAgravo.AUSENTE:
        return _SEM_AGRAVO
    return linha.agravo


def rotulo_exame(slug: str, exames_vocab: Mapping[str, Any], avisos: AvisosPCMSO) -> str:
    entrada = exames_vocab.get(slug, {})
    nome = str(entrada.get("nome_exibicao", slug))
    codigo = entrada.get("codigo_seconci")
    if not codigo:
        if nome not in avisos.exames_sem_codigo:
            avisos.exames_sem_codigo.append(nome)
        return nome
    return f"({codigo}) {nome}"


def celulas_do_exame(exame: ExameEmitido) -> tuple[str, ...]:
    """ADMISSÃO, APÓS ADM., PERIÓDICO, RET. TRAB, MUD. FUNÇ., DEMISSÃO. "Após adm." não existe no motor: sempre "-"."""
    def marca(momento: Momento) -> str:
        return "X" if momento in exame.momentos else "-"

    periodico = (
        f"{exame.periodicidade_meses} meses"
        if Momento.PER in exame.momentos and exame.periodicidade_meses > 0
        else "-"
    )
    return (marca(Momento.ADM), "-", periodico, marca(Momento.RT), marca(Momento.MR), marca(Momento.DEM))


# ----------------------------------------------------------------------------------------------- docx


def _bordas_e_fundo(celula: Any, fundo: str) -> None:
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement("w:tcBorders")
    for lado in ("top", "start", "bottom", "end"):
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:space"), "0")
        borda.set(qn("w:color"), "000000")
        bordas.append(borda)
    tc_pr.append(bordas)
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), fundo)
    tc_pr.append(sombra)


def _escrever_celula(celula: Any, texto: str, *, negrito: bool = False, centro: bool = False,
                     fundo: str = "FFFFFF") -> Paragraph:
    _bordas_e_fundo(celula, fundo)
    paragrafo: Paragraph = celula.paragraphs[0]
    if centro:
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragrafo.add_run(texto)
    run.bold = negrito
    return paragrafo


def _nova_tabela(doc: DocumentoDocx, linhas: int, larguras: Sequence[int]) -> Any:
    tabela = doc.add_table(rows=linhas, cols=len(larguras))
    tabela.autofit = False
    for coluna, largura in zip(tabela._tbl.tblGrid.findall(qn("w:gridCol")), larguras):
        coluna.set(qn("w:w"), str(largura))
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras):
            celula.width = Twips(largura)
    return tabela


def _campo_tc(paragrafo: Paragraph, titulo: str) -> None:
    """Entrada do índice (campo TC oculto, como no modelo)."""
    posicao = len(paragrafo._p.findall(qn("w:pPr")))
    partes = (("begin", None), (None, f' TC "{titulo}" \\l 1 '), ("separate", None), ("end", None))
    for deslocamento, (tipo, texto) in enumerate(partes):
        run = OxmlElement("w:r")
        if tipo is not None:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run.append(instr)
        paragrafo._p.insert(posicao + deslocamento, run)


def _bloco_ghe(doc: DocumentoDocx, matriz: MatrizGHE, ghe: GHEPGR | None, exames_vocab: Mapping[str, Any],
               vocabulario: VocabularioAgravos, avisos: AvisosPCMSO) -> list[Any]:
    elementos: list[Any] = []
    titulo = "Setor: " + _sanitizar(titulo_ghe_rq61(matriz.ghe_id, matriz.nome_ghe))
    setor = _nova_tabela(doc, 1, (_LARGURA_TOTAL,))
    paragrafo = _escrever_celula(setor.cell(0, 0), titulo, negrito=True, fundo=_FUNDO_CABECALHO)
    _campo_tc(paragrafo, titulo)
    elementos.append(setor._tbl)

    perigos = linhas_de_perigo(ghe, vocabulario, avisos) if ghe is not None else []
    if ghe is None:
        avisos.ghes_sem_pgr.append(matriz.ghe_id)
    cargos = [c for c in (_sanitizar(c) for c in matriz.cargos) if c]
    tabela = _nova_tabela(doc, len(cargos) + 1 + max(len(perigos), 1), _LARGURAS_PERIGOS)
    for i, cargo in enumerate(cargos):
        celula = tabela.cell(i, 0).merge(tabela.cell(i, 2))
        _escrever_celula(celula, f"Cargo: {cargo}", negrito=True)
    base = len(cargos)
    for coluna, rotulo in enumerate(_COLUNAS_PERIGOS):
        _escrever_celula(tabela.cell(base, coluna), rotulo, negrito=True, fundo=_FUNDO_CABECALHO)
    if not perigos:
        _escrever_celula(tabela.cell(base + 1, 0), "(perigos do PGR indisponíveis — conferir)")
        _escrever_celula(tabela.cell(base + 1, 1), "")
        _escrever_celula(tabela.cell(base + 1, 2), "")
    for i, linha in enumerate(perigos, start=base + 1):
        _escrever_celula(tabela.cell(i, 0), linha.perigo)
        _escrever_celula(tabela.cell(i, 1), linha.grupo)
        _escrever_celula(tabela.cell(i, 2), texto_do_agravo(linha))
    elementos.append(tabela._tbl)

    exames = sorted(matriz.linhas, key=lambda e: _chave_ordem_exame(e.exame, dict(exames_vocab)))
    tabela_exames = _nova_tabela(doc, len(exames) + 1, _LARGURAS_EXAMES)
    for coluna, rotulo in enumerate(_COLUNAS_EXAMES):
        _escrever_celula(tabela_exames.cell(0, coluna), rotulo, negrito=True, centro=coluna > 0,
                         fundo=_FUNDO_CABECALHO)
    for i, exame in enumerate(exames, start=1):
        _escrever_celula(tabela_exames.cell(i, 0), rotulo_exame(exame.exame, exames_vocab, avisos))
        for coluna, valor in enumerate(celulas_do_exame(exame), start=1):
            _escrever_celula(tabela_exames.cell(i, coluna), valor, centro=True)
    elementos.append(tabela_exames._tbl)

    for observacao in matriz.observacoes:
        elementos.append(doc.add_paragraph(_formatar_observacao(observacao, dict(exames_vocab)))._p)
    elementos.append(doc.add_paragraph()._p)
    return elementos


def _paragrafos_com_campos(doc: DocumentoDocx) -> list[Any]:
    partes = [doc.element.body]
    for secao in doc.sections:
        for cabecalho in (secao.header, secao.even_page_header, secao.first_page_header):
            partes.append(cabecalho._element)
    return [t for parte in partes for t in parte.iter(qn("w:t")) if t.text and "{{" in t.text]


def preencher_campos(doc: DocumentoDocx, valores: Mapping[str, str]) -> None:
    for t in _paragrafos_com_campos(doc):
        if t.text == _MARCA_GHE:
            continue
        t.text = _CAMPO.sub(lambda m: valores.get(m.group(1), EM_BRANCO), t.text)


def gerar_pcmso_docx(
    matrizes: Sequence[MatrizGHE],
    pgr: PGR | None,
    exames_vocab: Mapping[str, Any],
    vocabulario: VocabularioAgravos,
    dados: DadosPCMSO,
    destino: Path,
    modelo: Path = MODELO_PADRAO,
) -> AvisosPCMSO:
    """Grava o PCMSO em `destino` e devolve o que fica para completar ou conferir."""
    avisos = AvisosPCMSO()
    doc = Document(str(modelo))
    preencher_campos(doc, valores_dos_campos(dados, avisos))
    marca = [p for p in doc.element.body.iter(qn("w:p")) if "".join(t.text or "" for t in p.iter(qn("w:t"))) == _MARCA_GHE]
    if len(marca) != 1:
        raise ValueError(f"modelo do PCMSO com {len(marca)} marcas {_MARCA_GHE}, esperada 1")
    ghes_pgr = {g.id: g for g in pgr.ghes} if pgr is not None else {}
    for matriz in matrizes:
        for elemento in _bloco_ghe(doc, matriz, ghes_pgr.get(matriz.ghe_id), exames_vocab, vocabulario, avisos):
            marca[0].addprevious(elemento)
    marca[0].getparent().remove(marca[0])
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return avisos
