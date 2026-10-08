"""Memorial de raciocínio da matriz (D-ARQ-87 fatia 2) — anexo para as médicas,
separado da matriz assinada: exame a exame, por que foi pedido, com que
periodicidade, com que base e com que grau de certeza, em linguagem clínica
(`resumo_clinico` de cada regra). Apresentação-pura sobre `MatrizGHE` (D-ARQ-72):
lê o que o motor grava em `Motivo` (fatia 1) e em `Observacao`, sem decisão
clínica nova. O fundamento técnico de auditoria (`base_normativa`) fica fora
do documento — está no protocolo e na revisão da tela.
"""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from agente_medico.motor.estagios.gates import TIPO_PGR_SEM_INVENTARIO_PSICOSSOCIAL
from agente_medico.motor.estagios.pendencias_estruturais import (
    TIPO_CONTAMINANTE_A_CONFIRMAR,
    TIPO_MENOR_APRENDIZ_LISTA_TIP,
)
from agente_medico.motor.tipos import (
    GHEPGR,
    PGR,
    ExameEmitido,
    MatrizGHE,
    Motivo,
    Observacao,
    OrigemRisco,
    Pendencia,
)
from agente_medico.superficie.documento_matriz import (
    _COR_DESTAQUE,
    _COR_DESTAQUE_HEX,
    _COR_TEXTO_SOBRE_DESTAQUE,
    CabecalhoDocumento,
    _aplicar_fundo,
    _chave_ordem_exame,
    _formatar_celula,
    _formatar_momentos,
    _sanitizar,
    nome_ghe_exibicao,
    titulo_ghe,
)

# Grau de certeza, do mais forte ao mais fraco (convenções de status do
# PROTOCOLO, D-ARQ-22 Parte A), com o rótulo que a médica lê.
_CERTEZA: dict[str, int] = {"VALIDADO": 3, "DERIVADO": 2, "INTERPRETADO": 1}
ROTULO_CERTEZA: dict[int, str] = {
    3: "Protocolo validado pela coordenação",
    2: "Norma ou matrizes de referência",
    1: "Interpretação do sistema — confirmar",
    0: "Sem classificação — confirmar",
}

# Página paisagem com margens de 1,5 cm: ~24,9 cm úteis por tabela.
_COLUNAS_GHE: tuple[tuple[str, float], ...] = (
    ("Exame", 5.2),
    ("Por que foi pedido", 12.4),
    ("Base", 3.8),
    ("Correção", 3.5),
)
_COLUNAS_REGRAS: tuple[tuple[str, float], ...] = (
    ("Regra", 16.9),
    ("Base", 4.5),
    ("Correção", 3.5),
)
_COLUNAS_REVISAR: tuple[tuple[str, float], ...] = (
    ("Decisão", 11.4),
    ("Onde aparece", 10.0),
    ("Correção", 3.5),
)


@dataclass(frozen=True)
class LinhaMemorial:
    codigo: str
    exame: str
    porque: str
    certeza: str


@dataclass(frozen=True)
class DecisaoARevisar:
    """Uma regra interpretada (ou sem status) e onde ela decide sozinha: a
    médica revisa a decisão uma vez, não a mesma regra em cada GHE."""
    regra_id: str
    resumo: str
    onde: tuple[str, ...]
    codigos: tuple[str, ...]


@dataclass(frozen=True)
class BlocoMemorial:
    ghe_id: str
    nome_ghe: str
    linhas: tuple[LinhaMemorial, ...]
    nao_pedidos: tuple[str, ...]
    nao_reconhecidos: tuple[str, ...] = ()
    a_confirmar: tuple[str, ...] = ()
    agravos: tuple[str, ...] = ()
    agravos_pgr: tuple[str, ...] = ()
    alertas: tuple[str, ...] = ()


@dataclass(frozen=True)
class RegraUsada:
    regra_id: str
    resumo: str
    certeza: str


@dataclass(frozen=True)
class Memorial:
    revisar_primeiro: tuple[DecisaoARevisar, ...]
    blocos: tuple[BlocoMemorial, ...]
    contagem_certeza: tuple[tuple[str, int], ...]
    regras_usadas: tuple[RegraUsada, ...]
    avisos_pgr: tuple[str, ...] = ()


_LEGENDA_MOMENTOS = (
    "Momentos: ADM admissional; PER periódico (anual quando não traz meses); MRO mudança de riscos "
    "ocupacionais; RET retorno ao trabalho; DEM demissional."
)
_DATA_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def data_exibicao(texto: str) -> str:
    """Data ISO (aaaa-mm-dd) vira dd/mm/aaaa; qualquer outro texto sai como digitado."""
    iso = _DATA_ISO.fullmatch(texto.strip())
    return f"{iso[3]}/{iso[2]}/{iso[1]}" if iso else texto


_FIM_DE_FRASE = re.compile(r"(?<!Dra)(?<!Dr)\.\s+(?=[A-ZÁÉÍÓÚ])")


def primeira_frase(resumo: str) -> str:
    """O gatilho e o exame — a primeira frase do resumo clínico. Norma e origem
    da conduta ficam no resumo completo, uma vez só, em "Regras usadas"."""
    return _FIM_DE_FRASE.split(resumo, maxsplit=1)[0].rstrip(".") + "."


def resumos_do_protocolo(regras: Sequence[Mapping[str, Any]]) -> dict[str, str]:
    return {str(r["id"]): str(r["resumo_clinico"]) for r in regras if "resumo_clinico" in r}


def _nome_exame(slug: str, exames_vocab: dict[str, Any]) -> str:
    return str(exames_vocab.get(slug, {}).get("nome_exibicao", slug))


def codigo_linha(ghe_id: str, exame: str, exames_vocab: dict[str, Any]) -> str:
    """Código da linha para a correção voltar: GHE e exame pelo nome que a
    médica vê na matriz (D-ARQ-87 cl.4, fatia 3)."""
    return f"{ghe_id} · {_nome_exame(exame, exames_vocab)}"


def _nivel(status: str | None) -> int:
    return _CERTEZA.get(status or "", 0)


def _componentes(exame: ExameEmitido) -> list[list[Motivo]]:
    """O que decide a linha: as regras que pedem a periodicidade final e, para
    cada momento, as regras que o pedem."""
    define = [m for m in exame.motivos if m.periodicidade_meses == exame.periodicidade_meses]
    return [define] + [[m for m in exame.motivos if momento in m.momentos] for momento in exame.momentos]


def nivel_de_certeza(exame: ExameEmitido) -> int:
    """O elo mais fraco do que decide a linha: em cada componente vale a regra
    mais forte que o pede; a linha fica com o pior componente. Uma regra
    INTERPRETADO que só repete o que uma VALIDADO já pede não rebaixa a linha;
    a que é a única a pedir o DEM, rebaixa."""
    return min(max((_nivel(m.status_regra) for m in c), default=0) for c in _componentes(exame))


def regras_fracas(exame: ExameEmitido) -> tuple[str, ...]:
    """As regras de um componente cujo melhor apoio é interpretado ou sem status."""
    fracas: dict[str, None] = {}
    for componente in _componentes(exame):
        if max((_nivel(m.status_regra) for m in componente), default=0) <= _CERTEZA["INTERPRETADO"]:
            fracas.update(dict.fromkeys(m.regra_id for m in componente))
    return tuple(fracas)


def _por_regra(exame: ExameEmitido) -> dict[str, Motivo]:
    unicos: dict[str, Motivo] = {}
    for m in exame.motivos:
        unicos.setdefault(m.regra_id, m)
    return unicos


def _termo_exibicao(termo: str) -> str:
    """Termo do PGR como a médica lê (D-ARQ-88 fatia 2, Q2): o NUL que o PDF
    deixa no lugar de parênteses sai, e os espaços se juntam."""
    return " ".join(_sanitizar(termo).split())


def _agente_exibicao(slug: str) -> str:
    return slug.replace("_", " ")


def _descrever_origem(o: OrigemRisco) -> str:
    """D-ARQ-88 fatia 2: o agente pelo termo do PGR (Q1); risco de FDS já é
    nomeado pela própria fonte; sem termo, o slug legível."""
    if o.termo:
        texto = f"{_termo_exibicao(o.termo)} — {o.fonte}"
    elif o.fonte.startswith("FDS"):
        texto = o.fonte
    else:
        texto = f"{_agente_exibicao(o.agente)} — {o.fonte}"
    if o.presumida:
        # Q3: D-ARQ-68 cl.5 só presume por silêncio documental — vale para toda perna.
        texto += "; sem medição no PGR; pedido por precaução"
    return texto


def _origem(m: Motivo) -> str:
    return "; ".join(dict.fromkeys(_descrever_origem(o) for o in m.origens))


def _porque(exame: ExameEmitido, resumos: Mapping[str, str]) -> str:
    regras = _por_regra(exame)
    varias = len(regras) > 1
    partes: list[str] = []
    for m in regras.values():
        texto = primeira_frase(resumos.get(m.regra_id, "Regra sem resumo clínico — confirmar."))
        origem = _origem(m)
        if origem:
            texto += f" Origem: {origem}."
        if varias and m.periodicidade_meses is not None:
            pedido = ExameEmitido(
                exame=exame.exame, periodicidade_meses=m.periodicidade_meses, momentos=set(m.momentos)
            )
            texto += f" [pede: {_formatar_momentos(pedido, True)}]"
        partes.append(f"{texto} (ref. {m.regra_id})")
    if varias and len({m.periodicidade_meses for m in regras.values()}) > 1:
        partes.append(
            f"Mais de uma regra pede este exame: vale a menor periodicidade "
            f"({exame.periodicidade_meses} meses) e a soma dos momentos."
        )
    return "\n".join(partes)


def _linha(
    ghe_id: str, exame: ExameEmitido, exames_vocab: dict[str, Any], resumos: Mapping[str, str]
) -> LinhaMemorial:
    return LinhaMemorial(
        codigo=codigo_linha(ghe_id, exame.exame, exames_vocab),
        exame=_formatar_celula(exame, exames_vocab),
        porque=_sanitizar(_porque(exame, resumos)),
        certeza=ROTULO_CERTEZA[nivel_de_certeza(exame)],
    )


def _nao_pedido(obs: Observacao, exames_vocab: dict[str, Any]) -> str:
    exames = ", ".join(_nome_exame(slug, exames_vocab) for slug in obs.exames_dispensados)
    agente = "; ".join(dict.fromkeys(_termo_exibicao(t) for t in obs.termos)) or _agente_exibicao(obs.agente)
    motivo = f"{agente} com risco {obs.nivel_risco.lower()} no PGR"
    if obs.medicao is not None:
        motivo += f" e medição abaixo do nível de ação ({obs.medicao})"
    return _sanitizar(
        f"{exames}: não pedido — {motivo}. NR-07, item 7.5.12, b: o indicador depende da "
        f"classificação do risco no PGR; fica a menção no PCMSO. (ref. {obs.regra_dispensa})"
    )


_VIZINHO_RECUSADO = re.compile(r"aproximaria de '([a-z0-9_]+)'")


def _vizinho_recusado(ghe_id: str, termo: str, pendencias: Sequence[Pendencia]) -> str | None:
    for p in pendencias:
        if p.ghe_id == ghe_id and p.tipo == "fuzzy_recusado" and f"'{termo}'" in p.motivo:
            achado = _VIZINHO_RECUSADO.search(p.motivo)
            if achado:
                return achado[1]
    return None


def contaminantes_a_confirmar(matriz: MatrizGHE) -> tuple[str, ...]:
    """R-FDS-07: o pacote do contaminante não sai sem ele identificado; sem esta
    lista, a FDS que falta no GHE não aparece no anexo que a médica revisa."""
    return tuple(
        _sanitizar(f"{p.motivo}. (ref. {p.regra_origem})")
        for p in matriz.pendencias
        if p.tipo == TIPO_CONTAMINANTE_A_CONFIRMAR
    )


def alertas_do_ghe(matriz: MatrizGHE) -> tuple[str, ...]:
    """R-TIP-01: alerta para a revisão médica que não muda exame (menor aprendiz em GHE com
    risco da Lista TIP); sem esta lista ele ficaria só na pendência, fora do anexo."""
    return tuple(
        _sanitizar(f"{p.motivo}. (ref. {p.regra_origem})")
        for p in matriz.pendencias
        if p.tipo == TIPO_MENOR_APRENDIZ_LISTA_TIP
    )


def avisos_do_pgr(pendencias_globais: Sequence[Pendencia]) -> tuple[str, ...]:
    """R-PSY-07: aviso sobre o PGR inteiro (sem inventário psicossocial depois da vigência
    na NR-01), no resumo do memorial."""
    return tuple(
        _sanitizar(f"{p.motivo}. (ref. {p.regra_origem})")
        for p in pendencias_globais
        if p.tipo == TIPO_PGR_SEM_INVENTARIO_PSICOSSOCIAL
    )


def agravos_pgr_do_ghe(ghe: GHEPGR) -> tuple[str, ...]:
    """D-ARQ-93: agravo que o PGR escreve na linha de cada risco do GHE, verbatim — uma linha
    por risco com agravo lido ("Ruido: A exposição ao ruído…"), sem repetir a mesma."""
    linhas = (
        f"{_termo_exibicao(risco.termo or risco.agente or '')}: {_termo_exibicao(risco.agravo)}"
        for risco in ghe.riscos
        if risco.agravo
    )
    return tuple(dict.fromkeys(linhas))


def agravos_do_ghe(ghe: GHEPGR) -> tuple[str, ...]:
    """D-ARQ-92 fatia 3: agravos à saúde das FDS vinculadas ao GHE, uma linha por produto —
    o agravo em texto, como a FDS escreve, e o código H entre parênteses. FDS vinculada sem
    frase H de saúde legível também aparece, para conferir (D-ARQ-22)."""
    linhas: list[str] = []
    for produto in ghe.produtos_quimicos:
        if produto.fds is None:
            continue
        if not produto.fds.agravos:
            linhas.append(f"{produto.nome}: a FDS não traz frase H de saúde legível — conferir a seção 2.")
            continue
        partes = (
            f"{f.texto.rstrip('.')} ({f.codigo})" if f.texto else f"{f.codigo} (texto não legível na FDS)"
            for f in produto.fds.agravos
        )
        linhas.append(f"{produto.nome}: {'; '.join(partes)}.")
    return tuple(dict.fromkeys(_sanitizar(linha) for linha in linhas))


def riscos_nao_reconhecidos(ghe: GHEPGR, pendencias: Sequence[Pendencia]) -> tuple[str, ...]:
    """Riscos que o PGR declara no GHE e que não viraram agente: nenhum exame sai
    deles, e sem esta lista a lacuna não aparece no anexo que a médica revisa
    (T65, 30/09/2026: gesso no GHE 16 e PNOS com erro de digitação no GHE 18)."""
    linhas: list[str] = []
    for risco in ghe.riscos:
        if risco.agente is not None or not risco.termo:
            continue
        termo = _termo_exibicao(risco.termo)
        if risco.causa_nao_resolucao == "fracao_sem_agente":
            explicacao = "nomeia fração ou medida sem a substância; pedir a FDS ao elaborador do PGR (R-PGR-05)"
        elif risco.causa_nao_resolucao == "fuzzy_recusado":
            vizinho = _vizinho_recusado(ghe.id, risco.termo, pendencias)
            explicacao = (
                f"parece {_agente_exibicao(vizinho)}, mas grafia aproximada não é aceita para agente "
                "que dispara exame; conferir a grafia no PGR"
                if vizinho
                else "grafia aproximada de um agente, não aceita; conferir a grafia no PGR"
            )
        else:
            explicacao = "sem correspondência no vocabulário de agentes"
        linhas.append(_sanitizar(f"{termo}: {explicacao}."))
    return tuple(dict.fromkeys(linhas))


def _lista_ghes(ghes: Sequence[str], todos: set[str]) -> str:
    if len(todos) > 1 and set(ghes) == todos:
        return f"todos os GHEs ({len(todos)})"
    return ", ".join(ghes)


def _decisoes_a_revisar(
    matrizes: Sequence[MatrizGHE], exames_vocab: dict[str, Any], resumos: Mapping[str, str]
) -> tuple[DecisaoARevisar, ...]:
    todos = {m.ghe_id for m in matrizes}
    ghes_por_exame: dict[str, dict[str, list[str]]] = {}
    codigos: dict[str, list[str]] = {}
    for matriz in matrizes:
        for exame in matriz.linhas:
            for regra_id in regras_fracas(exame):
                ghes_por_exame.setdefault(regra_id, {}).setdefault(exame.exame, []).append(matriz.ghe_id)
                codigos.setdefault(regra_id, []).append(codigo_linha(matriz.ghe_id, exame.exame, exames_vocab))
    return tuple(
        DecisaoARevisar(
            regra_id=regra_id,
            resumo=resumos.get(regra_id, "Regra sem resumo clínico — confirmar."),
            onde=tuple(
                f"{_nome_exame(slug, exames_vocab)}: {_lista_ghes(ghes, todos)}" for slug, ghes in por_exame.items()
            ),
            codigos=tuple(codigos[regra_id]),
        )
        for regra_id, por_exame in ghes_por_exame.items()
    )


def montar_memorial(
    matrizes: Sequence[MatrizGHE],
    exames_vocab: dict[str, Any],
    resumos: Mapping[str, str],
    pgr: PGR | None = None,
    pendencias: Sequence[Pendencia] = (),
    pendencias_globais: Sequence[Pendencia] = (),
) -> Memorial:
    ghes_pgr = {g.id: g for g in pgr.ghes} if pgr is not None else {}
    blocos: list[BlocoMemorial] = []
    niveis: Counter[int] = Counter()
    usadas: dict[str, Motivo] = {}
    for matriz in matrizes:
        ordenadas = sorted(matriz.linhas, key=lambda e: _chave_ordem_exame(e.exame, exames_vocab))
        niveis.update(nivel_de_certeza(e) for e in ordenadas)
        for exame in ordenadas:
            for regra_id, motivo in _por_regra(exame).items():
                usadas.setdefault(regra_id, motivo)
        blocos.append(
            BlocoMemorial(
                ghe_id=matriz.ghe_id,
                nome_ghe=nome_ghe_exibicao(matriz.nome_ghe),
                linhas=tuple(_linha(matriz.ghe_id, e, exames_vocab, resumos) for e in ordenadas),
                nao_pedidos=tuple(_nao_pedido(o, exames_vocab) for o in matriz.observacoes),
                nao_reconhecidos=(
                    riscos_nao_reconhecidos(ghes_pgr[matriz.ghe_id], pendencias)
                    if matriz.ghe_id in ghes_pgr
                    else ()
                ),
                a_confirmar=contaminantes_a_confirmar(matriz),
                agravos=agravos_do_ghe(ghes_pgr[matriz.ghe_id]) if matriz.ghe_id in ghes_pgr else (),
                agravos_pgr=agravos_pgr_do_ghe(ghes_pgr[matriz.ghe_id]) if matriz.ghe_id in ghes_pgr else (),
                alertas=alertas_do_ghe(matriz),
            )
        )
    return Memorial(
        revisar_primeiro=_decisoes_a_revisar(matrizes, exames_vocab, resumos),
        blocos=tuple(blocos),
        contagem_certeza=tuple((ROTULO_CERTEZA[n], niveis[n]) for n in sorted(ROTULO_CERTEZA, reverse=True) if niveis[n]),
        regras_usadas=tuple(
            RegraUsada(
                regra_id=regra_id,
                resumo=_sanitizar(resumos.get(regra_id, "Regra sem resumo clínico — confirmar.")),
                certeza=ROTULO_CERTEZA[_nivel(motivo.status_regra)],
            )
            for regra_id, motivo in sorted(usadas.items())
        ),
        avisos_pgr=avisos_do_pgr(pendencias_globais),
    )


def linhas_da_tabela(bloco: BlocoMemorial) -> list[tuple[str, str, str, str]]:
    """Exames do mesmo GHE com o mesmo motivo e a mesma base viram uma linha só
    (ex.: hemograma, glicemia e ECG que vêm só da atividade crítica)."""
    grupos: dict[tuple[str, str], list[str]] = {}
    for linha in bloco.linhas:
        grupos.setdefault((linha.porque, linha.certeza), []).append(linha.exame)
    return [("\n".join(exames), porque, certeza, "") for (porque, certeza), exames in grupos.items()]


RODAPE_CONFIDENCIAL = "Uso interno — confidencial"
TITULO_AGRAVOS = "Agravos à saúde declarados nas FDS dos produtos deste GHE:"
TITULO_AGRAVOS_PGR = "Agravos à saúde declarados no PGR para os riscos deste GHE:"


def aplicar_rodape_confidencial(documento: Any) -> None:
    """Anexos para as médicas (memorial, riscos para o ASO) não vão ao cliente: o rodapé
    de cada página marca o sigilo. A matriz assinada não leva o rodapé."""
    for secao in documento.sections:
        paragrafo = secao.footer.paragraphs[0]
        paragrafo.text = RODAPE_CONFIDENCIAL
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.runs[0].font.size = Pt(8)


def _tabela(documento: Any, colunas: Sequence[tuple[str, float]], linhas: Sequence[Sequence[str]]) -> None:
    tabela = documento.add_table(rows=1, cols=len(colunas))
    tabela.style = "Table Grid"
    tabela.autofit = False
    for celula, (rotulo, _) in zip(tabela.rows[0].cells, colunas):
        celula.text = rotulo
        paragrafo = celula.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.runs[0].bold = True
        paragrafo.runs[0].font.color.rgb = _COR_TEXTO_SOBRE_DESTAQUE
        _aplicar_fundo(celula, _COR_DESTAQUE_HEX)
    for linha in linhas:
        for celula, texto in zip(tabela.add_row().cells, linha):
            celula.text = texto
            for paragrafo in celula.paragraphs:
                for run in paragrafo.runs:
                    run.font.size = Pt(9)
    # O LibreOffice lê a grade (gridCol) e o Word lê a célula: as duas levam a largura.
    for coluna, (_, largura) in zip(tabela.columns, colunas):
        coluna.width = Cm(largura)
    for linha_tabela in tabela.rows:
        for celula, (_, largura) in zip(linha_tabela.cells, colunas):
            celula.width = Cm(largura)


def renderizar_memorial_docx(memorial: Memorial, cabecalho: CabecalhoDocumento, destino: Path) -> None:
    documento = Document()
    for secao in documento.sections:
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = secao.page_height, secao.page_width
        secao.top_margin = secao.bottom_margin = Cm(1.5)
        secao.left_margin = secao.right_margin = Cm(1.5)
    aplicar_rodape_confidencial(documento)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("Memorial de raciocínio da matriz de exames")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = _COR_DESTAQUE
    documento.add_paragraph(
        f"Empresa: {cabecalho.empresa} · Obra: {cabecalho.obra} · Data: {data_exibicao(cabecalho.data)}"
    )
    documento.add_paragraph(
        "Anexo de conferência da matriz, não assinado. Para cada exame: por que foi pedido, com que "
        "base e com que grau de certeza. Quando mais de uma regra pede o mesmo exame, vale a menor "
        "periodicidade e a soma dos momentos. Nas tabelas por GHE aparece o motivo em uma frase; a norma e a "
        "origem de cada conduta estão na seção 3. Para corrigir, escreva na coluna Correção; a correção volta "
        "para o sistema e ajusta a regra citada em \"ref.\"."
    )
    documento.add_paragraph(_LEGENDA_MOMENTOS)

    total = sum(n for _, n in memorial.contagem_certeza)
    documento.add_heading("Resumo", level=1)
    documento.add_paragraph(f"{total} exames na matriz, em {len(memorial.blocos)} GHEs. Base de cada um:")
    for rotulo, n in memorial.contagem_certeza:
        documento.add_paragraph(f"{rotulo}: {n}", style="List Bullet")
    for aviso in memorial.avisos_pgr:
        documento.add_paragraph(aviso, style="List Bullet")
    nao_pedidos = sum(len(b.nao_pedidos) for b in memorial.blocos)
    if nao_pedidos:
        documento.add_paragraph(f"Exames dispensados pelo nível de risco do PGR: {nao_pedidos}", style="List Bullet")
    nao_reconhecidos = [b for b in memorial.blocos if b.nao_reconhecidos]
    if nao_reconhecidos:
        documento.add_paragraph(
            f"Riscos do PGR que o sistema não reconheceu — nenhum exame sai deles: "
            f"{sum(len(b.nao_reconhecidos) for b in nao_reconhecidos)}, em "
            f"{', '.join(b.ghe_id for b in nao_reconhecidos)}. Conferir antes de validar.",
            style="List Bullet",
        )
    a_confirmar = [b for b in memorial.blocos if b.a_confirmar]
    if a_confirmar:
        documento.add_paragraph(
            f"FDS a pedir ao elaborador do PGR para confirmar contaminante: "
            f"{', '.join(b.ghe_id for b in a_confirmar)}.",
            style="List Bullet",
        )
    com_alerta = [b for b in memorial.blocos if b.alertas]
    if com_alerta:
        documento.add_paragraph(
            f"Alertas para a revisão médica: {', '.join(b.ghe_id for b in com_alerta)}.",
            style="List Bullet",
        )

    documento.add_heading("1. Confirmar primeiro — decisões sem base direta em norma ou protocolo", level=1)
    if memorial.revisar_primeiro:
        _tabela(
            documento,
            _COLUNAS_REVISAR,
            [(f"{d.resumo} (ref. {d.regra_id})", "\n".join(d.onde), "") for d in memorial.revisar_primeiro],
        )
    else:
        documento.add_paragraph("Nenhum exame depende só de interpretação do sistema.")

    documento.add_heading("2. Matriz completa, por GHE", level=1)
    for bloco in memorial.blocos:
        cabecalho_ghe = documento.add_heading(titulo_ghe(bloco.ghe_id, bloco.nome_ghe), level=2)
        if cabecalho_ghe.runs:
            cabecalho_ghe.runs[0].font.color.rgb = _COR_DESTAQUE
        _tabela(documento, _COLUNAS_GHE, linhas_da_tabela(bloco))
        for texto in bloco.nao_pedidos:
            documento.add_paragraph(texto, style="List Bullet")
        if bloco.nao_reconhecidos:
            documento.add_paragraph("Riscos do PGR não reconhecidos (nenhum exame sai deles — conferir):")
            for texto in bloco.nao_reconhecidos:
                documento.add_paragraph(texto, style="List Bullet")
        if bloco.a_confirmar:
            documento.add_paragraph("FDS a pedir ao elaborador do PGR (contaminante a confirmar):")
            for texto in bloco.a_confirmar:
                documento.add_paragraph(texto, style="List Bullet")
        if bloco.alertas:
            documento.add_paragraph("Alertas para a revisão médica:")
            for texto in bloco.alertas:
                documento.add_paragraph(texto, style="List Bullet")
        if bloco.agravos_pgr:
            documento.add_paragraph(TITULO_AGRAVOS_PGR)
            for texto in bloco.agravos_pgr:
                documento.add_paragraph(texto, style="List Bullet")
        if bloco.agravos:
            documento.add_paragraph(TITULO_AGRAVOS)
            for texto in bloco.agravos:
                documento.add_paragraph(texto, style="List Bullet")

    documento.add_heading("3. Regras usadas nesta matriz — norma e origem de cada conduta", level=1)
    _tabela(
        documento,
        _COLUNAS_REGRAS,
        [(f"{r.resumo} (ref. {r.regra_id})", r.certeza, "") for r in memorial.regras_usadas],
    )
    documento.save(str(destino))
