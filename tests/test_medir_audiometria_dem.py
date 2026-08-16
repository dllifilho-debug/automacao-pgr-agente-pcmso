"""Testes de scripts/medir_audiometria_dem.py (003.EX fatia 0) — parsear_momentos
sobre células de gabarito isoladas à linha de Audiometria."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from agente_medico.motor.tipos import Momento
from agente_medico.superficie.documento_matriz import _ROTULO_MOMENTO
from scripts.medir_audiometria_dem import (
    RegistroCargo,
    classificar_dem,
    extrair_registros,
    parsear_momentos,
    rotulos_nao_reconhecidos,
)


def test_parsear_momentos_le_os_quatro() -> None:
    # Reversão que mata: fazer parsear_momentos retornar só o primeiro rótulo casado.
    resultado = parsear_momentos("Audiometria (ADM, PER, MRO, DEM)")
    assert resultado == {Momento.ADM, Momento.PER, Momento.MR, Momento.DEM}


def test_parsear_momentos_sem_dem_nao_inventa_dem() -> None:
    # Reversão que mata: devolver set(Momento) sempre que casar qualquer rótulo.
    resultado = parsear_momentos("Audiometria (ADM, PER, MRO)")
    assert Momento.DEM not in resultado
    assert resultado == {Momento.ADM, Momento.PER, Momento.MR}


def test_rotulo_desconhecido_nao_vira_momento_e_e_reportado() -> None:
    # Reversão que mata: engolir o desconhecido em silêncio (continue sem registrar) —
    # aplicada a rotulos_nao_reconhecidos, que passaria a sempre devolver frozenset() vazio.
    celula = "Audiometria (ADM, XPTO)"
    assert parsear_momentos(celula) == {Momento.ADM}
    assert rotulos_nao_reconhecidos(celula) == {"XPTO"}


def test_parsear_momentos_cobre_todo_par_da_inversao_computado_do_dado() -> None:
    # Reversão que mata: trocar _ROTULO_PARA_MOMENTO (a inversão importada de
    # _ROTULO_MOMENTO) por um dict literal redigitado no script, omitindo RET.
    for momento, rotulo in _ROTULO_MOMENTO.items():
        assert parsear_momentos(f"Audiometria ({rotulo})") == {momento}


def test_classificar_dem_separa_indeterminado_de_sem_dem_confirmado() -> None:
    # Reversão que mata: fazer a agregação somar célula com rótulo não
    # reconhecido em sem_dem (indeterminado sempre vazio) — dobra a recusa
    # do parser em adivinhar numa negação silenciosa na agregação.
    limpo = RegistroCargo(
        ghe="GHE-X",
        cargo="Cargo Limpo",
        tem_audiometria=True,
        momentos_audiometria=frozenset({Momento.ADM, Momento.PER, Momento.MR}),
        rotulos_nao_reconhecidos=frozenset(),
    )
    ambiguo = RegistroCargo(
        ghe="GHE-Y",
        cargo="Cargo Ambíguo",
        tem_audiometria=True,
        momentos_audiometria=frozenset({Momento.ADM, Momento.PER, Momento.MR}),
        rotulos_nao_reconhecidos=frozenset({"DEM 12 meses"}),
    )
    com_dem, indeterminado, sem_dem = classificar_dem([limpo, ambiguo])
    assert com_dem == []
    assert indeterminado == [ambiguo]
    assert sem_dem == [limpo]


def _construir_docx_duas_colunas(caminho: Path) -> None:
    documento = Document()
    tabela = documento.add_table(rows=0, cols=2)
    linha_cabecalho = tabela.add_row().cells
    linha_cabecalho[0].text = "FUNÇÃO"
    linha_cabecalho[1].text = "EXAMES SOLICITADOS"
    linha_cargo = tabela.add_row().cells
    linha_cargo[0].text = "Cargo Simples"
    linha_cargo[1].text = "Audiometria (ADM, PER, MRO, DEM)"
    documento.save(str(caminho))


def test_celulas_logicas_neutra_em_tabela_de_duas_colunas_sem_mesclagem(tmp_path: Path) -> None:
    # Controle de neutralidade da Entrega 3 (003.EZ, DH-003EY-01): em tabela
    # de 2 colunas físicas sem mesclagem, _celulas_logicas(linha.cells)
    # devolve exatamente [texto_col0, texto_col1] — o mesmo resultado da
    # indexação fixa celulas[0]/celulas[1] herdada de 003.EX. Nenhuma
    # reversão isolada mata este teste; ele confirma que o port não muda o
    # resultado nos documentos que já eram lidos corretamente (forma dos
    # checkpoints SPE 0030 e RESERVA 0028).
    caminho = tmp_path / "duas_colunas.docx"
    _construir_docx_duas_colunas(caminho)
    registros, _ = extrair_registros(caminho)
    assert len(registros) == 1
    assert registros[0].cargo == "Cargo Simples"
    assert registros[0].tem_audiometria is True
    assert registros[0].momentos_audiometria == {
        Momento.ADM,
        Momento.PER,
        Momento.MR,
        Momento.DEM,
    }


def _construir_docx_mesclado(caminho: Path) -> None:
    # python-docx repete o texto da célula mesclada em cada coluna física do
    # span — simulado aqui atribuindo o mesmo texto às colunas 0-1 (FUNÇÃO)
    # e 2-3 (EXAMES SOLICITADOS), sem chamar .merge() (o efeito sobre
    # .text é o mesmo que uma mesclagem real produziria na leitura).
    documento = Document()
    tabela = documento.add_table(rows=0, cols=4)
    linha_cabecalho = tabela.add_row().cells
    linha_cabecalho[0].text = "FUNÇÃO"
    linha_cabecalho[1].text = "FUNÇÃO"
    linha_cabecalho[2].text = "EXAMES SOLICITADOS"
    linha_cabecalho[3].text = "EXAMES SOLICITADOS"
    linha_cargo = tabela.add_row().cells
    linha_cargo[0].text = "Cargo Mesclado"
    linha_cargo[1].text = "Cargo Mesclado"
    linha_cargo[2].text = "Audiometria (ADM, PER, MRO, DEM)"
    linha_cargo[3].text = "Audiometria (ADM, PER, MRO, DEM)"
    documento.save(str(caminho))


def test_celulas_logicas_evita_zerar_audiometria_em_tabela_mesclada(tmp_path: Path) -> None:
    # Reversão que mata: reverter o port de _celulas_logicas (voltar a
    # indexação fixa celulas[0]/celulas[1]) — celulas[1] leria a própria
    # mesclagem de FUNÇÃO em vez de EXAMES SOLICITADOS, e tem_audiometria
    # cairia para False (reproduz o defeito medido no ATZUM, DH-003EY-01).
    caminho = tmp_path / "mesclado.docx"
    _construir_docx_mesclado(caminho)
    registros, _ = extrair_registros(caminho)
    assert len(registros) == 1
    assert registros[0].tem_audiometria is True
    assert registros[0].momentos_audiometria == {
        Momento.ADM,
        Momento.PER,
        Momento.MR,
        Momento.DEM,
    }
